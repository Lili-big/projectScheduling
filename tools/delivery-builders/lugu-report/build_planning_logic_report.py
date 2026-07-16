from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "docs" / "泸古项目上午调研验证记录_20260715.docx"
OUTPUT = ROOT / "docs" / "泸古项目前期工期策划思路分析_20260715.docx"

COLORS = {
    "navy": "1F4D78",
    "blue": "2E74B5",
    "ink": "202939",
    "muted": "667085",
    "light": "F2F4F7",
    "callout": "F4F6F9",
    "border": "D0D5DD",
    "yellow": "FFF4CC",
    "yellow_text": "7A5A00",
}
LATIN_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def clear_body(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_paragraph_border(paragraph, left=None, bottom=None, fill=None):
    p_pr = paragraph._p.get_or_add_pPr()
    if fill:
        shd = p_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            p_pr.append(shd)
        shd.set(qn("w:fill"), fill)
        shd.set(qn("w:val"), "clear")
    if left or bottom:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        if left:
            node = OxmlElement("w:left")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), str(left.get("size", 18)))
            node.set(qn("w:space"), str(left.get("space", 8)))
            node.set(qn("w:color"), left.get("color", COLORS["blue"]))
            borders.append(node)
        if bottom:
            node = OxmlElement("w:bottom")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), str(bottom.get("size", 12)))
            node.set(qn("w:space"), str(bottom.get("space", 8)))
            node.set(qn("w:color"), bottom.get("color", COLORS["blue"]))
            borders.append(node)


def add_callout(doc, label, text, tone="blue"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = True
    fill = COLORS["callout"] if tone == "blue" else COLORS["yellow"]
    border = COLORS["blue"] if tone == "blue" else COLORS["yellow_text"]
    add_paragraph_border(p, left={"color": border, "size": 22, "space": 10}, fill=fill)
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, size=11, bold=True, color=COLORS["navy"] if tone == "blue" else COLORS["yellow_text"])
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=COLORS["ink"])
    return p


def add_numbering(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    n_id = OxmlElement("w:numId")
    n_id.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(n_id)
    p_pr.append(num_pr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True, color=COLORS["ink"])
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11, color=COLORS["ink"])
    else:
        set_run_font(p.add_run(text), size=11, color=COLORS["ink"])
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True
    return p


def add_body(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True, color=COLORS["ink"])
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11, color=COLORS["ink"], italic=italic)
    else:
        set_run_font(p.add_run(text), size=11, color=COLORS["ink"], italic=italic)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    assert sum(widths_dxa) == 9360
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), COLORS["border"])
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    set_table_geometry(table, widths_dxa)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, COLORS["light"])
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2 if r_idx else 0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or c_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=9.0, bold=(r_idx == 0), color=COLORS["navy"] if r_idx == 0 else COLORS["ink"])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


doc = Document(TEMPLATE)
clear_body(doc)
doc.core_properties.title = "泸古项目前期工期策划思路分析"
doc.core_properties.subject = "客户前期总进度计划策划方法分析"
doc.core_properties.author = "广联达斑马产品团队"
doc.core_properties.keywords = "泸古项目, 工期策划, 总进度计划, 架梁倒排, 客户调研"

section = doc.sections[0]
header = section.header
hp = header.paragraphs[0]
hp.clear()
hp.paragraph_format.space_after = Pt(0)
hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
set_run_font(hp.add_run("客户工期策划方法分析"), size=8.5, bold=True, color=COLORS["muted"])
set_run_font(hp.add_run("\t泸古项目｜内部材料"), size=8.5, color=COLORS["muted"])

bullet_id = add_numbering(doc, "bullet")
decimal_id = add_numbering(doc, "decimal")

# 首页：standard_business_brief + memo_masthead
doc.add_paragraph().paragraph_format.space_after = Pt(12)
doc.add_paragraph("泸古项目前期工期策划思路分析", style="Title")
doc.add_paragraph("基于工程部长上午调研录音的业务方法还原", style="Subtitle")
metadata = [
    ("项目", "泸古项目"),
    ("访谈对象", "工程部长 李宗仁"),
    ("调研日期", "2026 年 7 月 15 日上午"),
    ("文档定位", "内部产品验证材料"),
    ("分析口径", "只归纳客户实际做法；我方介绍不作为客户已验证结论"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run(f"{label}："), size=10.5, bold=True, color="000000")
    set_run_font(p.add_run(value), size=10.5, color="000000")
rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(6)
rule.paragraph_format.space_after = Pt(10)
add_paragraph_border(rule, bottom={"color": COLORS["blue"], "size": 14, "space": 2})
add_callout(
    doc,
    "核心结论",
    "客户不是从项目开工日机械顺排全部任务，而是以架梁路线、连续梁合拢等控制节点搭建总控骨架，先按工程量、工效和资源计算持续时间，再倒推最晚开工时间；控制性工程精细排，普通工程保留现场调整空间。",
)

add_heading(doc, "一、客户前期工期策划的整体方法", 1)
add_body(doc, "客户的计划编制过程可以归纳为六个连续步骤。真正困难的部分不是在 Excel 中填日期，而是把架梁通道、控制节点、施工逻辑、资源周转和现场弹性先想清楚。")
steps = [
    "识别项目的合同工期、架梁、合拢等控制节点，明确哪些日期不能被突破。",
    "根据各工点工程量、经验工效和资源投入，计算每项工作的持续时间。",
    "以梁场和架梁路线串联全线工点，确定桥与桥之间的主要先后关系。",
    "从架梁或合拢节点向前倒排，形成各结构物的最迟完成和最晚开始时间。",
    "结合模板、钻机、挂篮、班组、钢筋加工和混凝土供应配置资源，并安排周转。",
    "按控制性工程与普通工程采用不同粒度，形成总体计划，并为月度滚动管理保留调整余量。",
]
for item in steps:
    add_list_item(doc, item, decimal_id)
add_callout(doc, "归纳表达", "任务持续时间≈工程量÷单资源工效÷有效资源数量＋养护、安装、转场及必要等待；最晚开始时间＝控制节点要求时间－任务持续时间－风险余量。")

add_heading(doc, "二、总体计划骨架：以架梁路线和控制节点为主线", 1)
add_heading(doc, "2.1 先确定梁场和架梁主线", 2)
for item in [
    "先明确梁场位置、建设和取证时间、开始制梁时间、初始存梁量及制梁能力。",
    "确定架桥机从哪里开始，按什么顺序经过各座桥，哪些桥必须优先贯通。",
    "考虑运梁车通道、架桥机横移、回退、拆装以及穿越隧道等条件。",
    "桥与桥之间最主要的逻辑关系来自架梁通道；普通下部结构在不影响架梁时可以灵活开工。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "2.2 连续刚构合拢决定通道释放", 2)
add_body(doc, "连续刚构并不是独立工点，而可能是架梁路线的制约点。只有主桥合拢并形成通道后，架桥机和运梁路线才能继续向后推进，因此连续梁的下部、0号块、挂篮和悬浇节段会成为全线控制链的一部分。")

add_heading(doc, "2.3 左右幅需要穿插，而非简单顺排", 2)
for item in [
    "不是始终完成一幅后再统一施工另一幅，而是结合架梁通道进行左右幅交叉。",
    "一幅架完后切换到另一幅，可给已架一幅留出横隔板、湿接缝等施工窗口。",
    "湿接缝混凝土施工后应减少运梁车持续通行产生的振动，避免影响结构耐久性。",
    "穿插方案的目标是减少等待时间，避免架梁期间其他工作面长期停滞并侵蚀总工期。",
]:
    add_list_item(doc, item, bullet_id)
add_callout(doc, "关键判断", "客户的架梁策划本质上是“通道释放＋设备移动＋结构施工窗口”的组合推演，不是按桥梁里程机械排序。", tone="yellow")

add_heading(doc, "三、工点持续时间与倒排方法", 1)
add_heading(doc, "3.1 按工程量、工效和资源测算持续时间", 2)
for item in [
    "桩基通常按桩数和钻机工效计算，例如一台钻机每天完成多少根。",
    "墩柱、柱系梁等结合结构数量、施工节奏和模板套数估算。",
    "连续梁按节段数量和单节段周期计算，并加入挂篮安装、预压等准备时间。",
    "任务之间还需要计入养护、转场、设备安装和必要等待。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "3.2 由最迟完成时间倒推出最晚开工", 2)
add_body(doc, "架梁开始时间确定后，客户会反推桥梁下部结构的最迟完成时间，并通常在架梁前预留约 1—2 个月余量。计划中的开始日期不是所有工点都必须严格执行的日期，而是判断是否仍有调整空间的基准。")
for item in [
    "可以提前施工，但不是所有任务越早完成越好。",
    "在可用余量内晚开或调整顺序，不视为真正风险。",
    "只有预测完成时间开始侵蚀架梁、合拢等控制节点余量时，才需要升级管理。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "四、连续梁与控制性工程的精细策划", 1)
add_body(doc, "客户对连续梁和关键通道工程会细化到主要工序，典型链路为：桩基→承台→墩身→0号块→挂篮安装及预压→悬浇节段→边跨合拢→中跨合拢。")
for item in [
    "比较多个 T 构的完成时间，以最晚完成的对象作为合拢控制点。",
    "其他 T 构不能提前过多，必要时主动推迟开工，使各 T 构尽量接近同步。",
    "控制最紧的主墩不能耽误，其他墩可以利用余量适当放慢。",
    "墩身、0号块和悬臂状态需要考虑连续施工与徐变差异，避免新老混凝土产生裂缝风险。",
    "边跨、过渡墩现浇段和中跨合拢需要按最终合拢时点协同安排。",
]:
    add_list_item(doc, item, bullet_id)
add_callout(doc, "策划特征", "客户追求的不是所有任务越早完成越好，而是关键结构在合适时间完成，并与最终合拢节点保持协调。")

add_heading(doc, "五、任务拆分粒度：重点细、普通粗", 1)
add_body(doc, "客户明确认为，计划过粗时 Excel 已经足够，软件没有额外价值；计划过细时又无法与现场实际施工结合。因此需要分层管理。")
add_table(
    doc,
    ["对象", "建议粒度", "现场允许的弹性"],
    [
        ["控制性工程", "细化到墩台、关键工序和控制节点", "关键墩台和先后关系需要强控制"],
        ["普通桥梁", "按桥、左右幅或月度完成量", "在总时间窗口内可调整具体施工顺序"],
        ["桩基", "按墩台或作业单元合并计算", "不固定每根桩的施工先后"],
        ["墩柱与柱系梁", "作为交替衔接的组合过程", "不能简单拆成两个互不关联的任务"],
        ["月度计划", "数量目标＋少量关键墩台要求", "普通任务允许现场选择具体对象"],
    ],
    [1800, 3360, 4200],
)
add_body(doc, "例如，同一墩台下多根桩基可以作为一个作业单元计算，现场可能先施工任意一根；只要在总体窗口内完成，并不要求严格按照计划编号顺序执行。")

add_heading(doc, "六、资源和工装策划：不是简单增加设备", 1)
add_heading(doc, "6.1 关键工装提前配置并安排周转", 2)
for item in [
    "控制性主墩通常单独配置模板，普通桥梁根据结构数量和工期确定模板套数。",
    "承台平钢模等通用模板可以周转，异形墩柱、柱系梁和盖梁模板可能需要单独配置。",
    "部分桥梁会有意后开，以等待其他桥梁模板周转过来。",
    "配置原则是既满足工期，又避免模板、设备和班组过量投入形成窝工。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "6.2 资源能力需要上下游耦合", 2)
add_body(doc, "钻机、模板或工作面增加，并不必然带来同比例工期压缩。客户会同时考虑钢筋加工、混凝土供应、班组能力和设备保障。")
for item in [
    "若钢筋加工厂每天只能供应一定数量钢筋笼，继续增加钻机只会形成等待。",
    "混凝土供应中断、设备故障或队伍不熟练会直接降低计划工效。",
    "梁场制梁、存梁和现场架梁速度必须匹配，不能只优化其中一个环节。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "6.3 日历与非理想因素", 2)
add_body(doc, "客户认为实际计划还应考虑高温、雨季、大风、春节、材料供应、设备故障和偶发地质问题等因素。经验上可按每月约 25 个有效施工日理解，并为全年保留一定非有效时间。")
add_callout(doc, "待确认口径", "录音中一方面提到真实计划应考虑这些因素，另一方面又提到当前版本尚未完整计入制约因素。需要由总工确认现有 Excel 已包含哪些余量、哪些仍是管理经验。", tone="yellow")

add_heading(doc, "七、现有 Excel 计划如何承载这套思路", 1)
for item in [
    "总体计划约有数百条任务，开始时间、持续时间和结束时间通过公式连接。",
    "修改关键输入后，后续日期可以自动顺延，支持总体计划快速调整。",
    "计划中还整理了架梁顺序、关键工点持续时间和主要工装配置。",
    "客户认为难点主要在前期思路和逻辑推演，横道图和公式只是承载结果。",
]:
    add_list_item(doc, item, bullet_id)
add_callout(doc, "产品边界", "如果系统只是复刻一张粗粒度横道图，客户认为 Excel 足够；系统必须在复杂逻辑复现、多方案重算、影响解释和月度滚动方面形成增量价值。")

add_heading(doc, "八、总体计划如何衔接月度执行", 1)
add_body(doc, "客户希望总体计划不是一次性成果，而是后续月度管理的基准。其现行业务闭环可归纳为：")
for item in [
    "从总体计划筛选当月应完成的任务和形象进度目标。",
    "按工点向施工队下达完成多少根桩、多少道系梁、多少米墩柱及少量关键墩台要求。",
    "由现场技术员核查实际完成情况，并记录未完成原因。",
    "未完成内容顺延到下月，同时要求采取增加资源或压缩周期等措施。",
    "通过月度进度分析会判断偏差是否已经影响总工期或控制节点。",
]:
    add_list_item(doc, item, decimal_id)
add_body(doc, "因此，前期总计划的价值不只在于得到一组日期，还要能够支持后续月度计划提取、实际进度回收、余量判断和滚动调整。")

add_heading(doc, "九、客户策划方法的核心原则", 1)
principles = [
    ["先骨架后细节", "先确定架梁、合拢和关键通道，再细化单体工程。"],
    ["先算持续时间再倒排", "先用工程量、工效和资源算工期，再根据最迟完成时间反推开工。"],
    ["关键位置强控制", "控制性工程细化到墩台和工序，普通工程控制总量和时间窗口。"],
    ["余量内允许调整", "普通日期偏差不是风险，侵蚀关键节点余量才需要预警。"],
    ["资源必须成套匹配", "钻机、模板、钢筋、混凝土、班组和工作面需要共同平衡。"],
    ["计划服务现场而非束缚现场", "硬约束必须执行，普通任务的具体顺序允许根据现场调整。"],
]
add_table(doc, ["原则", "具体含义"], principles, [2600, 6760])

add_heading(doc, "十、下一步验证建议", 1)
add_body(doc, "当前项目尚未开工，动态进度闭环暂不具备实测条件。下一步应优先验证客户的前期策划和计算过程，而不是只展示最终排程结果。")
for item in [
    "与总工复盘原始总计划的形成过程，获取工程量、工效、架梁顺序、资源配置和控制节点依据。",
    "选择“全线架梁控制链＋一座控制性桥梁”作为范围，在系统中复现客户现有计划。",
    "以开工调整到 9 月 1 日为真实变化条件，比较不调整、资源调整和顺序调整等方案。",
    "重点验证逻辑是否覆盖、计算是否可信、差异是否可解释，以及相比 Excel 是否减少重复推演时间。",
    "由总工评价结果是否能支持真实决策，而不只评价功能是否完整或页面是否好看。",
]:
    add_list_item(doc, item, decimal_id)

add_heading(doc, "十一、需要向总工继续确认的问题", 1)
questions = [
    "架梁路线确定时比较过哪些备选方案，最终方案胜出的核心原因是什么？",
    "开工日期调整到 9 月 1 日后，哪些控制节点仍固定，哪些任务或资源需要重新安排？",
    "关键工效来自定额、历史经验、分包承诺还是项目管理目标？",
    "现有 Excel 已计入哪些天气、节假日、材料供应和设备故障余量？",
    "哪些资源必须前置固定配置，哪些可以通过方案计算再决定？",
    "哪些施工逻辑属于不能违反的硬约束，哪些只是优选组织方式？",
    "客户花费时间最多的是数据整理、日期计算，还是方案和架梁路线推演？",
]
for item in questions:
    add_list_item(doc, item, decimal_id)

add_heading(doc, "附：资料来源与口径说明", 1)
add_body(doc, "资料来源为《泸古项目-工程部长-李宗仁调研-1.docx》和《泸古项目-工程部长-李宗仁调研-2.docx》。两份转写文件的说话人编号不同，分析时按文件首页角色映射识别客户表达。")
add_body(doc, "文档对明显转写错误进行了术语规范，例如“连续钢构”统一为“连续刚构”、“虚变”统一为“徐变”、“装气梁/柱西梁”按上下文统一为“柱系梁”等；无法确认的桥名和数字不作为核心结论依据。")
add_body(doc, "本报告中的计算公式为对客户方法的产品化归纳，不代表客户提供了正式算法或已确认全部参数口径。", italic=True)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
