from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


WORKPACKAGE = Path(__file__).resolve().parents[1]
PATH = WORKPACKAGE / "results" / "泸古项目7月16日前期工期策划思路分析_20260715.docx"

with ZipFile(PATH) as archive:
    bad = archive.testzip()
    assert bad is None, bad

doc = Document(PATH)
assert len(doc.sections) == 1
section = doc.sections[0]
assert round(section.page_width.inches, 2) == 8.50
assert round(section.page_height.inches, 2) == 11.00
assert all(
    round(value.inches, 2) == 1.00
    for value in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin)
)

paragraphs = [p for p in doc.paragraphs if p.text.strip()]
headings = [p for p in paragraphs if p.style.name.startswith("Heading")]
numbered = [p for p in paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
all_text = "\n".join(p.text for p in paragraphs)

assert paragraphs[0].text == "泸古项目前期工期策划思路分析"
assert len(headings) >= 20
assert len(numbered) >= 50
assert len(doc.tables) == 2
assert "说话人1" not in all_text
assert "说话人2" not in all_text
assert "00:" not in all_text
assert "PLACEHOLDER" not in all_text
content_text = "\n".join(p.text for p in paragraphs[:-3])
assert "虚变" not in content_text
assert "连续钢构" not in content_text

for table in doc.tables:
    tbl_pr = table._tbl.tblPr
    widths = [int(x.get(qn("w:w"))) for x in table._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert sum(widths) == 9360, widths
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    layout = tbl_pr.find(qn("w:tblLayout"))
    assert tbl_w is not None and tbl_w.get(qn("w:w")) == "9360"
    assert tbl_ind is not None and tbl_ind.get(qn("w:w")) == "120"
    assert layout is not None and layout.get(qn("w:type")) == "fixed"
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            assert tc_w is not None and int(tc_w.get(qn("w:w"))) == widths[idx]
            assert cell.text.strip()

print(
    {
        "path": str(PATH),
        "bytes": PATH.stat().st_size,
        "paragraphs": len(paragraphs),
        "headings": len(headings),
        "real_lists": len(numbered),
        "tables": len(doc.tables),
        "zip_ok": True,
        "geometry_ok": True,
    }
)
