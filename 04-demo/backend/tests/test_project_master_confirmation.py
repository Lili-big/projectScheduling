from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook


BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from app.contracts.project_master import ConfirmProjectMasterVersionRequest  # noqa: E402
from app.project_master.repository import ProjectMasterConflictError, ProjectMasterRepository  # noqa: E402
from app.project_master.service import ProjectMasterService  # noqa: E402
from project_master_fixture_helpers import valid_project_master_workbook  # noqa: E402


def test_confirm_is_optimistic_atomic_and_history_is_immutable(tmp_path: Path) -> None:
    repository = ProjectMasterRepository(tmp_path / "master.db")
    invalidated: list[str] = []
    service = ProjectMasterService(repository, version_invalidation_handler=invalidated.append)
    first = service.import_workbook(
        project_id="demo",
        file_name="first.xlsx",
        content=valid_project_master_workbook(),
        created_by="tester",
        expected_current_version_id=None,
    )
    first_id = first.created_version_id
    assert first_id
    service.confirm_version(first_id, ConfirmProjectMasterVersionRequest(confirmed_by="reviewer"))

    second = service.import_workbook(
        project_id="demo",
        file_name="second.xlsx",
        content=valid_project_master_workbook(workpoint_name="第二版"),
        created_by="tester",
        expected_current_version_id=first_id,
    )
    second_id = second.created_version_id
    assert second_id
    with pytest.raises(ProjectMasterConflictError) as error:
        service.confirm_version(
            second_id,
            ConfirmProjectMasterVersionRequest(confirmed_by="reviewer", expected_current_version_id="stale"),
        )
    assert error.value.code == "CURRENT_VERSION_CHANGED"
    assert repository.get_version_summary(first_id).status == "confirmed"
    assert repository.get_version_summary(second_id).status == "draft"

    service.confirm_version(
        second_id,
        ConfirmProjectMasterVersionRequest(confirmed_by="reviewer", expected_current_version_id=first_id),
    )
    assert repository.get_version_summary(first_id).status == "superseded"
    assert repository.get_version_summary(second_id).status == "confirmed"
    assert invalidated == [first_id]
    with pytest.raises(ProjectMasterConflictError) as immutable:
        repository.confirm_version(second_id, expected_current_version_id=second_id, confirmed_by="again")
    assert immutable.value.code == "VERSION_IMMUTABLE"


def test_warnings_require_acknowledgement(tmp_path: Path) -> None:
    repository = ProjectMasterRepository(tmp_path / "master.db")
    service = ProjectMasterService(repository)
    batch = service.import_workbook(
        project_id="demo",
        file_name="warning.xlsx",
        content=valid_project_master_workbook(zero_quantity=True),
        created_by="tester",
        expected_current_version_id=None,
    )
    assert batch.created_version_id and batch.counts.warnings == 1
    with pytest.raises(ProjectMasterConflictError) as warning:
        service.confirm_version(
            batch.created_version_id,
            ConfirmProjectMasterVersionRequest(confirmed_by="reviewer"),
        )
    assert warning.value.code == "WARNING_NOT_ACKNOWLEDGED"
    confirmed = service.confirm_version(
        batch.created_version_id,
        ConfirmProjectMasterVersionRequest(
            confirmed_by="reviewer",
            acknowledge_warning_codes=["ZERO_QUANTITY"],
        ),
    )
    assert confirmed.status == "confirmed"


def test_referenced_deletion_blocks_confirmation(tmp_path: Path) -> None:
    repository = ProjectMasterRepository(tmp_path / "master.db")
    first_service = ProjectMasterService(repository)
    first = first_service.import_workbook(
        project_id="demo",
        file_name="first.xlsx",
        content=valid_project_master_workbook(),
        created_by="tester",
        expected_current_version_id=None,
    )
    first_id = first.created_version_id or ""
    first_service.confirm_version(first_id, ConfirmProjectMasterVersionRequest(confirmed_by="reviewer"))

    workbook = load_workbook(BytesIO(valid_project_master_workbook()))
    workpoints = workbook["工点信息"]
    workpoints.delete_rows(next(row for row in range(3, workpoints.max_row + 1) if workpoints.cell(row, 1).value == "WP-T01"))
    structures = workbook["结构物信息"]
    structures.delete_rows(next(row for row in range(3, structures.max_row + 1) if structures.cell(row, 1).value == "ST-TN-01"))
    output = BytesIO()
    workbook.save(output)
    service = ProjectMasterService(repository, reference_conflict_checker=lambda _version, deleted: ["实绩引用"] if "WP-T01" in deleted else [])
    second = service.import_workbook(
        project_id="demo",
        file_name="second.xlsx",
        content=output.getvalue(),
        created_by="tester",
        expected_current_version_id=first_id,
    )
    assert second.created_version_id
    with pytest.raises(ProjectMasterConflictError) as conflict:
        service.confirm_version(
            second.created_version_id,
            ConfirmProjectMasterVersionRequest(confirmed_by="reviewer", expected_current_version_id=first_id),
        )
    assert conflict.value.code == "REFERENCE_CONFLICT"
