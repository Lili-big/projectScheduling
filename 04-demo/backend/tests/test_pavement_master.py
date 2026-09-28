import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.contracts.project_master import ProjectMasterSnapshot
from app.project_master.workbook import create_template_bytes, export_snapshot, parse_workbook
from app.project_master.validation import validate_snapshot


def test_handover_three_states_and_legacy_roundtrip():
    from app.project_master.validation import resolve_roadbed_handover
    from app.contracts.project_master import ParameterValue
    import pytest
    assert resolve_roadbed_handover({"roadbed_available_date": "2026-10-25"})[0] == "dated"
    assert resolve_roadbed_handover({})[0] == "pending"
    snapshot = pavement_snapshot()
    for section, status in zip(snapshot.workpoints[0].structures, ["dated", "handed_over", "pending", "pending"]):
        if status != "dated":
            section.parameters = [p for p in section.parameters if p.parameter_code != "roadbed_available_date"]
        section.parameters.extend([ParameterValue(parameter_code="roadbed_handover_status", value_type="text", value=status),
            ParameterValue(parameter_code="roadbed_handover_note", value_type="text", value="征地未解决" if status == "pending" else "")])
    restored, issues, _ = parse_workbook(export_snapshot(snapshot))
    assert not [i for i in [*issues, *validate_snapshot(restored)] if i.severity == "error"]
    assert [resolve_roadbed_handover({p.parameter_code:p.value for p in s.parameters})[0] for s in restored.workpoints[0].structures] == ["dated", "handed_over", "pending", "pending"]
    for props in ({"roadbed_handover_status":"dated"}, {"roadbed_handover_status":"bad"},
                  {"roadbed_available_date":"bad"}, {"roadbed_handover_status":"pending", "roadbed_available_date":"2026-10-25"}):
        with pytest.raises(ValueError): resolve_roadbed_handover(props)


def pavement_snapshot():
    return ProjectMasterSnapshot.model_validate({"workpoints": [{
        "workpoint_id": "ROAD", "workpoint_name": "演示路面", "workpoint_type": "pavement", "schedule_support": "pavement_supported",
        "structures": [{
            "structure_id": f"{segment}-{side}", "workpoint_id": "ROAD", "structure_name": f"{segment}段{side}",
            "structure_category": "pavement", "structure_type": "pavement_section", "side": side,
            "parameters": [{"parameter_code": k, "value_type": typ, "value": v} for k, typ, v in [
                ("start_chainage", "text", "K0+000" if segment == "A" else "K1+000"),
                ("end_chainage", "text", "K0+800" if segment == "A" else "K1+500"),
                ("construction_length_m", "number", 760 if segment == "A" else 410), ("width_m", "number", 10),
                ("water_stable_thickness_m", "number", 0.76), ("water_stable_density_t_m3", "number", 2.38),
                ("roadbed_available_date", "date", "2026-01-01"), ("quantity_basis_confirmed", "boolean", True),
                ("quantity_basis_note", "text", "扣除桥涵后的确认净长")]],
            "components": [{"component_id": f"{segment}-{side}-W", "structure_id": f"{segment}-{side}",
                "component_name": "水稳底基层", "component_type": "cement_stabilized_base", "quantity": 760 if segment == "A" else 410,
                "unit": "m", "sort_order": 1, "parameters": [
                    {"parameter_code": "thickness_m", "value_type": "number", "value": 0.2},
                    {"parameter_code": "quantity_basis", "value_type": "text", "value": "entered"}]}]
        } for segment in ["A", "B"] for side in ["left", "right"]]
    }]})


def test_pending_execution_gaps_remain_savable_but_are_scheduling_diagnostics():
    from app.contracts.project_master import ParameterValue
    from app.project_master.validation import roadbed_start_offset
    from datetime import date
    import pytest
    snapshot = pavement_snapshot()
    section = snapshot.workpoints[0].structures[0]
    section.parameters = [p for p in section.parameters if p.parameter_code != "roadbed_available_date"]
    section.parameters.append(ParameterValue(parameter_code="roadbed_handover_status", value_type="text", value="pending"))
    section.components[0].parameters = []
    issues = validate_snapshot(snapshot)
    assert any(i.issue_code == "PAVEMENT_DATA_INCOMPLETE" and i.object_id == section.components[0].component_id for i in issues)
    assert not [i for i in issues if i.severity == "error"]
    with pytest.raises(ValueError): roadbed_start_offset({"roadbed_handover_status": "pending"}, date(2026, 9, 23))
    assert roadbed_start_offset({"roadbed_handover_status": "pending"}, date(2026, 9, 23), allow_pending=True) == 0


def test_disabled_section_is_distinct_from_missing_layers():
    snapshot = pavement_snapshot()
    section = snapshot.workpoints[0].structures[0]
    for component in section.components:
        component.enabled = False
    assert not [i for i in validate_snapshot(snapshot) if i.object_id == section.structure_id]
    section.components = []
    assert any(i.issue_code == "PAVEMENT_DATA_INCOMPLETE" and i.object_id == section.structure_id
               for i in validate_snapshot(snapshot))


@pytest.mark.parametrize("thickness", [None, "", 0, -0.1, "unknown", float("inf")])
def test_length_scheduling_does_not_depend_on_thickness(thickness):
    from app.project_master.validation import pavement_quantity_errors
    props = {"quantity_basis": "geometric", "quantity_basis_confirmed": True,
             "quantity_basis_note": "confirmed net length", "construction_length_m": 1001,
             "width_m": 10, "thickness_m": thickness}
    assert pavement_quantity_errors(props, 1001, "m") == []
    assert any(code == "PAVEMENT_QUANTITY_BASIS_UNCONFIRMED"
               for code, _ in pavement_quantity_errors(props, 1000, "m"))
    for unit in ("m2", "m3", "t"):
        assert any("thickness_m" in message for _, message in pavement_quantity_errors(props, 1001, unit))
    props.pop("thickness_m")
    assert pavement_quantity_errors(props, 1001, "m") == []


@pytest.mark.parametrize("pending", [False, True])
def test_master_length_scheduling_ignores_thickness_for_all_handover_states(pending):
    from app.contracts.project_master import ParameterValue
    snapshot = pavement_snapshot()
    section = snapshot.workpoints[0].structures[0]
    if pending:
        section.parameters = [p for p in section.parameters if p.parameter_code != "roadbed_available_date"]
        section.parameters.append(ParameterValue(parameter_code="roadbed_handover_status", value_type="text", value="pending"))
    for parameter in section.components[0].parameters:
        if parameter.parameter_code == "thickness_m":
            parameter.value = 0
    assert not [issue for issue in validate_snapshot(snapshot) if "thickness_m" in issue.message]


def test_pavement_roundtrip_preserves_net_length_and_source():
    snapshot = pavement_snapshot()
    restored, issues, fingerprint = parse_workbook(export_snapshot(snapshot))
    assert not [x for x in [*issues, *validate_snapshot(restored)] if x.severity == "error"]
    assert len(restored.workpoints[0].structures) == 4
    section = restored.workpoints[0].structures[0]
    assert {p.parameter_code: p.value for p in section.parameters}["construction_length_m"] == 760
    assert {p.parameter_code: p.value for p in section.parameters}["water_stable_thickness_m"] == 0.76
    assert {p.parameter_code: p.value for p in section.parameters}["water_stable_density_t_m3"] == 2.38
    assert section.components[0].parameters[0].parameter_code == "thickness_m"
    assert section.components[0].parameters[0].value == 0.2
    assert section.source.row_no == 3
    assert parse_workbook(export_snapshot(restored))[2] == fingerprint


def test_template_domain_and_legacy_versions():
    book = load_workbook(BytesIO(create_template_bytes("pavement")))
    assert book["填写说明"].cell(1, 2).value == "1.2"
    assert "param.construction_length_m" in [c.value for c in book["结构物信息"][1]]
    for version in ["1.0", "1.1"]:
        book["填写说明"].cell(1, 2).value = version
        out = BytesIO(); book.save(out)
        assert "TEMPLATE_VERSION_UNSUPPORTED" not in {i.issue_code for i in parse_workbook(out.getvalue())[1]}


def test_missing_planning_data_is_warning_but_duplicate_identity_blocks():
    snapshot = pavement_snapshot()
    snapshot.workpoints[0].structures[0].parameters = []
    issues = validate_snapshot(snapshot)
    assert any(i.issue_code == "PAVEMENT_DATA_INCOMPLETE" and i.severity == "warning" for i in issues)
    snapshot.workpoints[0].structures[1].structure_id = snapshot.workpoints[0].structures[0].structure_id
    assert any(i.issue_code == "STABLE_ID_DUPLICATE" and i.severity == "error" for i in validate_snapshot(snapshot))


def test_bad_unit_and_ambiguous_layer_order_are_visible():
    snapshot = pavement_snapshot()
    section = snapshot.workpoints[0].structures[0]
    section.components[0].unit = "个"
    clone = section.components[0].model_copy(update={"component_id": "extra"})
    section.components.append(clone)
    codes = {i.issue_code for i in validate_snapshot(snapshot)}
    assert "PAVEMENT_UNIT_MISMATCH" in codes
    assert "PAVEMENT_LAYER_ORDER_INVALID" in codes
