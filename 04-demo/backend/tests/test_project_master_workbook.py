from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.project_master.validation import validate_snapshot  # noqa: E402
from app.project_master.workbook import (  # noqa: E402
    create_template_bytes,
    export_snapshot,
    parse_workbook,
)
from project_master_fixture_helpers import (  # noqa: E402
    invalid_project_master_workbook,
    valid_project_master_workbook,
)


def test_template_has_fixed_four_sheets_and_two_header_rows() -> None:
    workbook = load_workbook(BytesIO(create_template_bytes()), data_only=True)
    assert workbook.sheetnames == ["填写说明", "工点信息", "结构物信息", "构件参数"]
    assert workbook["工点信息"].cell(1, 1).value == "workpoint_id"
    assert workbook["工点信息"].cell(2, 1).value == "工点ID"


def test_parse_valid_workbook_keeps_one_bridge_and_left_right_structures() -> None:
    snapshot, parse_issues, fingerprint = parse_workbook(valid_project_master_workbook())
    issues = [*parse_issues, *validate_snapshot(snapshot)]

    assert not [item for item in issues if item.severity == "error"]
    assert len(snapshot.workpoints) == 3
    bridge = next(item for item in snapshot.workpoints if item.workpoint_type == "bridge")
    assert {item.side for item in bridge.structures} >= {"left", "right", "shared"}
    assert len(fingerprint) == 64


def test_invalid_parent_and_case_insensitive_duplicate_are_blocking() -> None:
    snapshot, parse_issues, _ = parse_workbook(invalid_project_master_workbook())
    issues = [*parse_issues, *validate_snapshot(snapshot)]
    codes = {item.issue_code for item in issues if item.severity == "error"}

    assert "PARENT_WORKPOINT_NOT_FOUND" in codes
    assert "STABLE_ID_DUPLICATE" in codes


def test_exported_snapshot_can_be_imported_without_business_loss() -> None:
    original, _, original_fingerprint = parse_workbook(valid_project_master_workbook())
    exported = export_snapshot(original)
    restored, issues, restored_fingerprint = parse_workbook(exported)

    assert not [item for item in issues if item.severity == "error"]
    assert restored_fingerprint == original_fingerprint


def test_fast_reader_keeps_side_and_non_negative_quantity_validation() -> None:
    workbook = load_workbook(BytesIO(valid_project_master_workbook()))
    workbook["结构物信息"].cell(3, 6).value = "middle"
    workbook["构件参数"].cell(3, 5).value = -1
    output = BytesIO()
    workbook.save(output)

    snapshot, parse_issues, _ = parse_workbook(output.getvalue())
    issues = [*parse_issues, *validate_snapshot(snapshot)]
    codes = {item.issue_code for item in issues if item.severity == "error"}

    assert "STRUCTURE_SIDE_INVALID" in codes
    assert "FIELD_TYPE_INVALID" in codes


def test_abutment_body_and_cap_beam_workbook_round_trip_keep_shared_validation() -> None:
    workbook = load_workbook(BytesIO(valid_project_master_workbook()))
    components = workbook["构件参数"]
    components.append(
        ["CP-AB-BODY", "ST-S-A0", "来源构件甲", "abutment_body", 1, "个", "是", 2, None, None, None, 6.5]
    )
    components.append(
        ["CP-PIER-CAP", "ST-L-P1", "来源构件乙", "cap_beam", 1, "个", "是", 2, None, None, None, 12]
    )
    components.append(
        ["CP-AB-DISABLED", "ST-S-A0", "来源构件丙", "abutment_body", 1, "个", "否", 3]
    )
    output = BytesIO()
    workbook.save(output)

    snapshot, parse_issues, fingerprint = parse_workbook(output.getvalue())
    issues = [*parse_issues, *validate_snapshot(snapshot)]
    assert not [item for item in issues if item.severity == "error"]

    by_id = {
        component.component_id: component
        for workpoint in snapshot.workpoints
        for structure in workpoint.structures
        for component in structure.components
    }
    assert by_id["CP-AB-BODY"].component_type == "abutment_body"
    assert by_id["CP-PIER-CAP"].component_type == "cap_beam"
    assert by_id["CP-AB-DISABLED"].enabled is False

    restored, restored_issues, restored_fingerprint = parse_workbook(export_snapshot(snapshot))
    assert not [item for item in [*restored_issues, *validate_snapshot(restored)] if item.severity == "error"]
    assert restored_fingerprint == fingerprint
    restored_disabled = next(
        component
        for workpoint in restored.workpoints
        for structure in workpoint.structures
        for component in structure.components
        if component.component_id == "CP-AB-DISABLED"
    )
    assert restored_disabled.enabled is False


def test_abutment_workbook_invalid_identity_and_quantity_use_generic_errors() -> None:
    workbook = load_workbook(BytesIO(valid_project_master_workbook()))
    components = workbook["构件参数"]
    components.append([None, "ST-S-A0", "缺少标识构件", "abutment_body", 1, "个", "是", 2])
    components.append(["CP-AB-NEG", "ST-S-A0", "无效数量构件", "abutment_body", -1, "个", "是", 3])
    output = BytesIO()
    workbook.save(output)

    snapshot, parse_issues, _ = parse_workbook(output.getvalue())
    issues = [*parse_issues, *validate_snapshot(snapshot)]
    codes = {item.issue_code for item in issues if item.severity == "error"}

    assert {"REQUIRED_FIELD", "FIELD_TYPE_INVALID"} <= codes
    component_ids = {
        component.component_id
        for workpoint in snapshot.workpoints
        for structure in workpoint.structures
        for component in structure.components
    }
    assert "CP-AB-NEG" not in component_ids
