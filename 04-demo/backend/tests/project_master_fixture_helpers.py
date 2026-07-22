from __future__ import annotations

from copy import deepcopy
import json
import sys
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from app.contracts.project_master import ProjectMasterSnapshot, ProjectMasterVersionSummary  # noqa: E402
from app.project_master.workbook import create_template_bytes  # noqa: E402


ABUTMENT_PROJECTION_BASELINE = (
    Path(__file__).parent / "fixtures" / "project_master" / "abutment-projection-baseline.json"
)


def load_abutment_projection_baseline() -> dict[str, Any]:
    return deepcopy(json.loads(ABUTMENT_PROJECTION_BASELINE.read_text(encoding="utf-8")))


def abutment_projection_snapshot() -> ProjectMasterSnapshot:
    payload = load_abutment_projection_baseline()
    return ProjectMasterSnapshot.model_validate(payload["snapshot"]).model_copy(deep=True)


def confirmed_abutment_version_summary() -> ProjectMasterVersionSummary:
    payload = load_abutment_projection_baseline()
    return ProjectMasterVersionSummary.model_validate(payload["version"]).model_copy(deep=True)


def confirmed_abutment_projection() -> tuple[ProjectMasterVersionSummary, ProjectMasterSnapshot]:
    return confirmed_abutment_version_summary(), abutment_projection_snapshot()


def valid_project_master_workbook(*, workpoint_name: str = "一号特大桥", zero_quantity: bool = False) -> bytes:
    workbook = load_workbook(BytesIO(create_template_bytes()))
    workpoints = workbook["工点信息"]
    workpoints.append(["WP-B01", workpoint_name, "bridge", "Z1", 1000, 1800, 1, None])
    workpoints.append(["WP-R01", "一区段路基", "roadbed", "Z1", 1800, 2400, 2, None])
    workpoints.append(["WP-T01", "一号隧道", "tunnel", "Z1", 2400, 3600, 3, None])

    structures = workbook["结构物信息"]
    structures.append(["ST-L-P1", "WP-B01", "左幅1号墩", "substructure", "bridge_pier", "left", "WS-L", "左幅工区", "key", 1, None])
    structures.append(["ST-R-P1", "WP-B01", "右幅1号墩", "substructure", "bridge_pier", "right", "WS-R", "右幅工区", "key", 2, None])
    structures.append(["ST-L-P2", "WP-B01", "左幅2号墩", "substructure", "bridge_pier", "left", "WS-L", "左幅工区", "normal", 3, None])
    structures.append(["ST-R-P2", "WP-B01", "右幅2号墩", "substructure", "bridge_pier", "right", "WS-R", "右幅工区", "control", 4, None])
    structures.append(["ST-S-A0", "WP-B01", "共用0号台", "substructure", "bridge_abutment", "shared", "WS-C", "共用工区", None, 0, None])
    structures.append(["ST-L-SP1", "WP-B01", "左幅第1跨", "superstructure", "simple_span", "left", "WS-L", "左幅工区", None, 10, None, 1, 40, "ST-S-A0", "ST-L-P1", None, None, None, 10])
    structures.append(["ST-R-SP1", "WP-B01", "右幅第1跨", "superstructure", "simple_span", "right", "WS-R", "右幅工区", None, 11, None, 1, 40, "ST-S-A0", "ST-R-P1", None, None, None, 10])
    structures.append(["ST-L-CIP2", "WP-B01", "左幅第2联现浇梁", "superstructure", "cast_in_place_unit", "left", "WS-L", "左幅工区", "key", 12, None, 2, 60, "ST-L-P1", "ST-L-P2", "3x20m", None, None, None])
    structures.append(["ST-R-CB2", "WP-B01", "右幅第2联连续梁", "superstructure", "continuous_unit", "right", "WS-R", "右幅工区", "control", 13, None, 2, 80, "ST-R-P1", "ST-R-P2", "40+40m", "ST-R-P1,ST-R-P2", 12, None])
    structures.append(["ST-RD-01", "WP-R01", "路基主体", "earthwork", "roadbed_section", "none", None, None, None, 1, None])
    structures.append(["ST-TN-01", "WP-T01", "隧道洞身", "tunnel_body", "tunnel_body", "none", None, None, None, 1, None])

    components = workbook["构件参数"]
    components.append(["CP-L-P1-PILE", "ST-L-P1", "左幅1号墩桩基", "pile", 0 if zero_quantity else 8, "根", "是", 1, None, 2.0, 45, None, "钻孔桩"])
    components.append(["CP-R-P1-PILE", "ST-R-P1", "右幅1号墩桩基", "pile", 8, "根", "是", 1, None, 2.0, 45, None, "钻孔桩"])
    components.append(["CP-L-SP1-BEAM", "ST-L-SP1", "左幅第1跨预制梁", "precast_beam", 10, "片", "是", 1, None])
    components.append(["CP-R-SP1-BEAM", "ST-R-SP1", "右幅第1跨预制梁", "precast_beam", 10, "片", "是", 1, None])
    components.append(["CP-L-CIP2", "ST-L-CIP2", "左幅第2联现浇箱梁", "cast_in_place_box_beam", 1, "联", "是", 1, None])
    components.append(["CP-R-CB2", "ST-R-CB2", "右幅第2联连续梁节段", "cast_in_place_continuous_beam", 12, "段", "是", 1, None])
    placements = workbook["线路关系"]
    placements.append(["RP-B01-L", "WP-B01", "left", "ZK", 1000, 1800, "SG-001", 1])
    placements.append(["RP-B01-R", "WP-B01", "right", "K", 1010, 1810, "SG-001", 1])
    placements.append(["RP-R01-L", "WP-R01", "left", "BK", 0, 600, "SG-002", 2])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def invalid_project_master_workbook() -> bytes:
    workbook = load_workbook(BytesIO(valid_project_master_workbook()))
    structures = workbook["结构物信息"]
    structures.append(["ST-ORPHAN", "WP-MISSING", "孤立结构", "substructure", "bridge_pier", "none", "WS-X", "未知工区", None, 99])
    workpoints = workbook["工点信息"]
    workpoints.append(["wp-b01", "重复桥梁", "bridge", "Z1", 1000, 1800, 9])
    placements = workbook["线路关系"]
    placements.append(["RP-B01-L-DUP", "WP-B01", "left", "ZK", 1000, 900, "SG-001", 2])
    placements.append(["RP-ORPHAN", "WP-MISSING", "right", "AK", 0, 100, "SG-X", 3])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def large_project_master_workbook(
    *,
    workpoint_count: int = 500,
    structure_count: int = 10_000,
    component_count: int = 50_000,
) -> bytes:
    workbook = load_workbook(BytesIO(create_template_bytes()))
    workpoints = workbook["工点信息"]
    structures = workbook["结构物信息"]
    components = workbook["构件参数"]
    for index in range(workpoint_count):
        workpoints.append([f"WP-R{index:04d}", f"路基工点{index:04d}", "roadbed", "Z1", index * 100, index * 100 + 99, index, None])
    for index in range(structure_count):
        workpoint_index = index % workpoint_count
        structures.append([
            f"ST-R{index:05d}",
            f"WP-R{workpoint_index:04d}",
            f"路基结构{index:05d}",
            "earthwork",
            "roadbed_section",
            "none",
            None,
            None,
            None,
            index,
            None,
        ])
    for index in range(component_count):
        structure_index = index % structure_count
        components.append([
            f"CP-R{index:06d}",
            f"ST-R{structure_index:05d}",
            f"构件{index:06d}",
            "other",
            1,
            "个",
            "是",
            index,
            None,
        ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
