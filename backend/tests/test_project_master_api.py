from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.bootstrap import create_app  # noqa: E402
from app.project_master.repository import ProjectMasterRepository  # noqa: E402
from app.project_master.repository import ProjectMasterRepositoryError  # noqa: E402
from app.project_master.service import ProjectMasterService  # noqa: E402
from asgi_client import request  # noqa: E402
from project_master_fixture_helpers import invalid_project_master_workbook, valid_project_master_workbook  # noqa: E402


def _app(tmp_path: Path, *, max_bytes: int = 25 * 1024 * 1024):
    app = create_app()
    app.state.project_master_service = ProjectMasterService(
        ProjectMasterRepository(tmp_path / "project-master.db"),
        import_max_bytes=max_bytes,
    )
    return app


def _multipart(fields: dict[str, str], content: bytes, file_name: str = "master.xlsx") -> tuple[bytes, str]:
    boundary = "----project-master-contract"
    chunks: list[bytes] = []
    for name, value in fields.items():
        chunks.extend(
            [
                f"--{boundary}\r\n".encode(),
                f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
                value.encode("utf-8"),
                b"\r\n",
            ]
        )
    chunks.extend(
        [
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="file"; filename="{file_name}"\r\n'.encode(),
            b"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n",
            content,
            b"\r\n",
            f"--{boundary}--\r\n".encode(),
        ]
    )
    return b"".join(chunks), f"multipart/form-data; boundary={boundary}"


def _upload(app, workbook: bytes, **fields: str):
    body, content_type = _multipart({"created_by": "tester", **fields}, workbook)
    return request(
        app,
        "POST",
        "/api/projects/demo/project-master/imports",
        raw_body=body,
        extra_headers={"content-type": content_type},
    )


def _json(body: bytes):
    return json.loads(body.decode("utf-8"))


def test_template_import_batch_query_and_cancel(tmp_path: Path) -> None:
    app = _app(tmp_path)
    status, headers, body = request(app, "GET", "/api/project-master/template")
    assert status == 200
    assert "spreadsheetml.sheet" in headers["content-type"]
    assert load_workbook(filename=__import__("io").BytesIO(body)).sheetnames == [
        "填写说明",
        "工点信息",
        "结构物信息",
        "构件参数",
    ]

    status, _, body = _upload(app, valid_project_master_workbook())
    batch = _json(body)
    assert status == 201
    assert batch["status"] == "ready"
    assert batch["counts"]["workpoints"] >= 2

    status, _, body = request(app, "GET", f"/api/project-master/imports/{batch['batch_id']}")
    assert status == 200 and _json(body)["created_version_id"] == batch["created_version_id"]

    status, _, body = request(
        app,
        "POST",
        f"/api/project-master/imports/{batch['batch_id']}",
        {"cancelled_by": "tester", "cancel_reason": "重新填写"},
    )
    assert status == 200
    cancelled = _json(body)
    assert cancelled["status"] == "cancelled"
    assert cancelled["created_version_id"] is None
    assert cancelled["cancelled_version_id"] == batch["created_version_id"]


def test_blocked_unchanged_conflict_and_payload_limit(tmp_path: Path) -> None:
    app = _app(tmp_path)
    status, _, body = _upload(app, invalid_project_master_workbook())
    assert status == 202
    assert _json(body)["status"] == "blocked"
    assert _json(body)["counts"]["errors"] > 0

    status, _, body = _upload(app, valid_project_master_workbook())
    draft = _json(body)
    version_id = draft["created_version_id"]
    status, _, body = request(
        app,
        "POST",
        f"/api/project-master/versions/{version_id}/confirm",
        {"confirmed_by": "reviewer", "expected_current_version_id": None, "acknowledge_warning_codes": []},
    )
    assert status == 200 and _json(body)["status"] == "confirmed"

    status, _, body = _upload(app, valid_project_master_workbook(), expected_current_version_id=version_id)
    assert status == 202
    assert _json(body)["status"] == "unchanged"
    assert _json(body)["existing_version_id"] == version_id

    modified = valid_project_master_workbook(workpoint_name="修改后的桥梁")
    status, _, body = _upload(app, modified, expected_current_version_id="stale")
    assert status == 409
    assert _json(body)["detail"]["code"] == "CURRENT_VERSION_CHANGED"

    small_app = _app(tmp_path / "small", max_bytes=10)
    status, _, body = _upload(small_app, valid_project_master_workbook())
    assert status == 413
    assert _json(body)["detail"]["code"] == "PAYLOAD_TOO_LARGE"

    invalid_app = _app(tmp_path / "invalid")
    status, _, body = _upload(invalid_app, b"not-an-excel-file")
    assert status == 202
    assert _json(body)["status"] == "blocked"
    assert _json(body)["issues"][0]["issue_code"] == "WORKBOOK_INVALID"


def test_repository_unavailable_maps_to_stable_503(tmp_path: Path) -> None:
    app = _app(tmp_path)

    class UnavailableService:
        def template_bytes(self):
            raise ProjectMasterRepositoryError("数据库不可用")

    app.state.project_master_service = UnavailableService()
    status, _, body = request(app, "GET", "/api/project-master/template")
    assert status == 503
    assert _json(body)["detail"]["code"] == "PROJECT_MASTER_STORAGE_ERROR"
