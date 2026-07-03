from __future__ import annotations

import base64
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from ..models import AiParameterMaterialKind, AiParameterUploadedMaterialSummary, ValidationMessage


MAX_AI_PARAMETER_FILE_COUNT = 10
MAX_AI_PARAMETER_TOTAL_BYTES = 50 * 1024 * 1024
SUPPORTED_SUFFIX_KIND: dict[str, AiParameterMaterialKind] = {
    ".txt": "text",
    ".md": "text",
    ".csv": "excel",
    ".tsv": "excel",
    ".xlsx": "excel",
    ".xlsm": "excel",
    ".docx": "word",
    ".doc": "word",
    ".pdf": "pdf",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
    ".bmp": "image",
}


class AiParameterMaterialError(ValueError):
    def __init__(self, message: str, *, status_code: int = 400) -> None:
        super().__init__(message)
        self.status_code = status_code


@dataclass(frozen=True)
class AiParameterMaterial:
    material_id: str
    file_name: str
    kind: AiParameterMaterialKind
    size_bytes: int
    content_text: str
    source_summary: str
    parse_status: str
    mime_type: str | None = None
    content_base64: str | None = None
    error_message: str | None = None

    def summary(self) -> AiParameterUploadedMaterialSummary:
        return AiParameterUploadedMaterialSummary(
            material_id=self.material_id,
            file_name=self.file_name,
            kind=self.kind,
            size_bytes=self.size_bytes,
            parse_status=self.parse_status,  # type: ignore[arg-type]
            source_summary=self.source_summary,
            error_message=self.error_message,
        )

    def ai_payload(self) -> dict[str, Any]:
        return {
            "material_id": self.material_id,
            "file_name": self.file_name,
            "kind": self.kind,
            "size_bytes": self.size_bytes,
            "mime_type": self.mime_type,
            "content_text": self.content_text[:20_000],
            "content_base64": self.content_base64,
            "source_summary": self.source_summary,
            "parse_status": self.parse_status,
            "error_message": self.error_message,
        }


def collect_ai_parameter_materials(
    fields: dict[str, str],
    files: dict[str, dict[str, bytes | str]],
) -> tuple[list[AiParameterMaterial], list[ValidationMessage]]:
    raw_texts = _text_inputs(fields)
    raw_files = _ordered_files(files)
    if len(raw_files) > MAX_AI_PARAMETER_FILE_COUNT:
        raise AiParameterMaterialError(f"单次最多上传 {MAX_AI_PARAMETER_FILE_COUNT} 个文件，请拆分资料后重试。")
    total_size = sum(len(item.get("content", b"") if isinstance(item.get("content"), bytes) else b"") for item in raw_files)
    if total_size > MAX_AI_PARAMETER_TOTAL_BYTES:
        raise AiParameterMaterialError("单次上传文件总大小不能超过 50MB，请拆分资料后重试。")
    if not raw_texts and not raw_files:
        raise AiParameterMaterialError("请至少提供一段文本或一个资料文件。")

    materials: list[AiParameterMaterial] = []
    warnings: list[ValidationMessage] = []
    for index, text in enumerate(raw_texts, start=1):
        text = text.strip()
        if not text:
            continue
        materials.append(
            AiParameterMaterial(
                material_id=f"mat_text_{index:03d}",
                file_name=f"粘贴文本 {index}",
                kind="text",
                size_bytes=len(text.encode("utf-8")),
                content_text=text,
                source_summary=_summary_from_text(text, "用户粘贴文本"),
                parse_status="parsed",
            )
        )

    for index, item in enumerate(raw_files, start=1):
        file_name = str(item.get("filename") or f"uploaded-{index}")
        content = item.get("content", b"")
        if not isinstance(content, bytes):
            content = bytes(str(content), "utf-8")
        try:
            material = _material_from_file(index, file_name, content)
        except AiParameterMaterialError as exc:
            warnings.append(ValidationMessage(level="warning", message=str(exc), subject_id=file_name))
            material = AiParameterMaterial(
                material_id=f"mat_file_{index:03d}",
                file_name=file_name,
                kind=_kind_from_name(file_name) or "text",
                size_bytes=len(content),
                content_text="",
                source_summary=f"{file_name} 解析失败",
                parse_status="failed",
                error_message=str(exc),
            )
        materials.append(material)

    if not materials:
        raise AiParameterMaterialError("提供的资料为空，请补充可解析文本或文件。")
    return materials, warnings


def _text_inputs(fields: dict[str, str]) -> list[str]:
    values: list[str] = []
    for key, value in fields.items():
        if key in {"text", "text_input", "text_inputs", "text_inputs[]"} or key.startswith("text_input_"):
            values.append(value)
    return values


def _ordered_files(files: dict[str, dict[str, bytes | str]]) -> list[dict[str, bytes | str]]:
    preferred = ["file", "files", "files[]"]
    ordered: list[dict[str, bytes | str]] = []
    seen: set[str] = set()
    for key in preferred:
        if key in files:
            ordered.append(files[key])
            seen.add(key)
    for key in sorted(files):
        if key not in seen:
            ordered.append(files[key])
    return ordered


def _material_from_file(index: int, file_name: str, content: bytes) -> AiParameterMaterial:
    if not content:
        raise AiParameterMaterialError(f"{file_name} 文件为空。", status_code=422)
    kind = _kind_from_name(file_name)
    if kind is None:
        raise AiParameterMaterialError(f"暂不支持 {Path(file_name).suffix or '无扩展名'} 文件。", status_code=415)
    if kind == "excel":
        content_text, status, error = _extract_excel_text(file_name, content)
    elif kind == "text":
        content_text, status, error = _decode_text(content), "parsed", None
    elif kind == "image":
        content_text, status, error = "", "partially_parsed", None
    else:
        content_text = _decode_text(content)
        status = "partially_parsed" if content_text else "partially_parsed"
        error = None if content_text else "当前未启用本地全文解析，已把文件元数据提交给 AI adapter。"
    summary = _summary_from_text(content_text, file_name) if content_text else f"{file_name}（{kind}，{len(content)} bytes）"
    return AiParameterMaterial(
        material_id=f"mat_file_{index:03d}",
        file_name=file_name,
        kind=kind,
        size_bytes=len(content),
        content_text=content_text,
        source_summary=summary,
        parse_status=status,
        mime_type=_mime_type_for(file_name, kind),
        content_base64=_binary_payload_for(kind, content),
        error_message=error,
    )


def _kind_from_name(file_name: str) -> AiParameterMaterialKind | None:
    return SUPPORTED_SUFFIX_KIND.get(Path(file_name).suffix.lower())


def _mime_type_for(file_name: str, kind: AiParameterMaterialKind) -> str | None:
    suffix = Path(file_name).suffix.lower()
    if kind == "image":
        return {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".webp": "image/webp",
            ".bmp": "image/bmp",
        }.get(suffix)
    if suffix == ".pdf":
        return "application/pdf"
    if suffix == ".docx":
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if suffix == ".doc":
        return "application/msword"
    return None


def _binary_payload_for(kind: AiParameterMaterialKind, content: bytes) -> str | None:
    if kind in {"word", "pdf", "image"}:
        return base64.b64encode(content).decode("ascii")
    return None


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore").strip()


def _extract_excel_text(file_name: str, content: bytes) -> tuple[str, str, str | None]:
    if Path(file_name).suffix.lower() in {".csv", ".tsv"}:
        return _decode_text(content), "parsed", None
    try:
        from io import BytesIO

        from openpyxl import load_workbook
    except ImportError:
        return "", "partially_parsed", "缺少 openpyxl，已把文件元数据提交给 AI adapter。"
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        return "", "failed", f"Excel 读取失败：{exc}"

    chunks: list[str] = []
    for worksheet in workbook.worksheets[:5]:
        chunks.append(f"[Sheet: {worksheet.title}]")
        row_count = 0
        for row in worksheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if not values:
                continue
            chunks.append(" | ".join(values[:12]))
            row_count += 1
            if row_count >= 80:
                break
    text = "\n".join(chunks).strip()
    return text, "parsed" if text else "failed", None if text else "Excel 未识别到有效文本。"


def _summary_from_text(text: str, source: str) -> str:
    compact = " ".join(text.split())
    if len(compact) > 120:
        compact = compact[:117] + "..."
    return f"{source}: {compact}" if compact else f"{source}: 无可展示文本"
