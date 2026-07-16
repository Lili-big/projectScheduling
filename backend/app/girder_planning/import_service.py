from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from ..contracts import (
    FieldCandidateValue,
    FieldConflict,
    GirderImportPreview,
    GirderWorkPoint,
    PassageConditionRef,
    ProjectDataVersion,
    SourceEvidence,
    ValidationMessage,
)
from .fingerprints import stable_id


class GirderImportError(ValueError):
    status_code = 422


_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("工点名称", "名称", "桥名", "name", "workpoint_name"),
    "workpoint_type": ("工点类型", "类型", "type", "workpoint_type"),
    "side": ("幅别", "左右幅", "side"),
    "mileage_start_m": ("起点里程", "里程起点", "起始里程", "mileage_start_m", "start_mileage"),
    "mileage_end_m": ("终点里程", "里程终点", "结束里程", "mileage_end_m", "end_mileage"),
    "corridor_id": ("走廊", "线路", "通道", "corridor_id"),
    "bridge_id": ("桥梁ID", "桥梁id", "bridge_id"),
    "work_section_id": ("工区ID", "工区id", "work_section_id"),
    "requires_erection": ("是否架梁", "待架", "requires_erection"),
    "explicit_readiness_date": ("明确开放日期", "通行日期", "readiness_date"),
    "linked_condition_refs": ("关联条件", "关联工程", "linked_condition_refs"),
}


def import_workpoints(
    *,
    file_name: str,
    content: bytes,
    project_version: ProjectDataVersion,
    coarse_mode: bool = False,
) -> GirderImportPreview:
    rows, sheet_name = _read_rows(file_name, content)
    if not rows:
        raise GirderImportError("架梁工点文件没有可读取的数据行。")
    header_index, headers = _find_headers(rows)
    field_indexes = _field_indexes(headers)
    if "name" not in field_indexes:
        raise GirderImportError("架梁工点文件缺少“工点名称/桥名”列。")

    workpoints: list[GirderWorkPoint] = []
    evidence: list[SourceEvidence] = []
    conflicts: list[FieldConflict] = []
    diagnostics: list[ValidationMessage] = []
    bridges_by_id = {item.id: item for item in project_version.project.bridges}
    bridges_by_name = {_normalize_text(item.name): item for item in project_version.project.bridges}

    for row_no, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = {field: _cell(row, index) for field, index in field_indexes.items()}
        if not any(value not in {None, ""} for value in values.values()):
            continue
        name = str(values.get("name") or "").strip()
        if not name:
            diagnostics.append(_message("warning", "GIRDER_IMPORT_EMPTY_NAME", f"第 {row_no} 行缺少工点名称，已跳过。"))
            continue
        workpoint_type = _workpoint_type(values.get("workpoint_type"), name)
        side = _side(values.get("side"))
        bridge = None
        imported_bridge_id = str(values.get("bridge_id") or "").strip()
        if workpoint_type == "bridge":
            bridge = bridges_by_id.get(imported_bridge_id) if imported_bridge_id else None
            bridge = bridge or bridges_by_name.get(_normalize_text(name))
            if bridge is None:
                conflict_id = stable_id("field-conflict", {"row": row_no, "field": "bridge_id", "value": imported_bridge_id or name})
                evidence_id = _append_evidence(
                    evidence,
                    file_name=file_name,
                    sheet_name=sheet_name,
                    row_no=row_no,
                    field_path=f"workpoints[{row_no}].bridge_id",
                    original_value=imported_bridge_id or name,
                    normalized_value=None,
                    authority_domain="girder_workpoint",
                )
                conflicts.append(
                    FieldConflict(
                        conflict_id=conflict_id,
                        entity_ref=f"row:{row_no}",
                        field_path="bridge_id",
                        severity="blocking",
                        candidate_values=[FieldCandidateValue(source_evidence_id=evidence_id, value=imported_bridge_id or name)],
                    )
                )
                diagnostics.append(_message("error", "GIRDER_BRIDGE_UNMAPPED", f"第 {row_no} 行桥梁“{name}”无法映射到项目结构。", f"row:{row_no}"))
                continue

        candidates = _expand_bridge_sides(bridge, side, coarse_mode) if bridge else [(None, side)]
        for section, expanded_side in candidates:
            bridge_id = bridge.id if bridge else None
            work_section_id = section.id if section else str(values.get("work_section_id") or "").strip() or None
            rough = expanded_side == "unknown" or bool(coarse_mode and workpoint_type == "bridge" and section is None)
            if workpoint_type == "bridge" and expanded_side == "both" and not coarse_mode:
                diagnostics.append(_message("error", "GIRDER_SIDE_BOTH_UNEXPANDED", f"桥梁“{name}”的幅别 both 无法展开为左右幅。", bridge_id or name))
            mileage_start = _mileage(values.get("mileage_start_m"))
            mileage_end = _mileage(values.get("mileage_end_m"))
            if mileage_end < mileage_start:
                mileage_start, mileage_end = mileage_end, mileage_start
            corridor_id = str(values.get("corridor_id") or "main").strip() or "main"
            suffix = work_section_id or expanded_side
            workpoint_id = (
                f"bridge:{bridge_id}:{suffix}"
                if bridge_id
                else stable_id("workpoint", {"name": name, "side": expanded_side, "mileage": [mileage_start, mileage_end], "row": row_no})
            )
            requires_erection = _boolean(values.get("requires_erection"), default=workpoint_type == "bridge")
            linked_refs = _linked_refs(values.get("linked_condition_refs"))
            readiness_date = _date_value(values.get("explicit_readiness_date"))
            workpoint = GirderWorkPoint(
                workpoint_id=workpoint_id,
                name=bridge.name if bridge else name,
                workpoint_type=workpoint_type,
                side=expanded_side,
                mileage_start_m=mileage_start,
                mileage_end_m=mileage_end,
                corridor_id=corridor_id,
                bridge_id=bridge_id,
                work_section_id=work_section_id,
                requires_erection=requires_erection,
                rough_granularity=rough,
                explicit_readiness_date=readiness_date,
                linked_condition_refs=linked_refs,
                properties={"source_row": row_no},
            )
            workpoints.append(workpoint)
            for field, normalized in (
                ("name", workpoint.name),
                ("side", workpoint.side),
                ("mileage_start_m", mileage_start),
                ("mileage_end_m", mileage_end),
                ("corridor_id", corridor_id),
                ("bridge_id", bridge_id),
                ("work_section_id", work_section_id),
            ):
                _append_evidence(
                    evidence,
                    file_name=file_name,
                    sheet_name=sheet_name,
                    row_no=row_no,
                    field_path=f"workpoints[{workpoint_id}].{field}",
                    original_value=values.get(field),
                    normalized_value=normalized,
                    authority_domain="structure" if field in {"bridge_id", "work_section_id"} and bridge else "girder_workpoint",
                )
            if rough:
                diagnostics.append(_message("warning", "GIRDER_ROUGH_GRANULARITY", f"工点“{workpoint.name}”仅可用于粗粒度预览。", workpoint_id))

    deduped: list[GirderWorkPoint] = []
    seen: set[str] = set()
    for item in workpoints:
        if item.workpoint_id in seen:
            diagnostics.append(_message("warning", "GIRDER_DUPLICATE_WORKPOINT", f"工点 {item.workpoint_id} 重复，已合并为一个稳定对象。", item.workpoint_id))
            continue
        seen.add(item.workpoint_id)
        deduped.append(item)
    return GirderImportPreview(
        workpoints=deduped,
        source_evidence=evidence,
        field_conflicts=conflicts,
        diagnostics=diagnostics,
    )


def _read_rows(file_name: str, content: bytes) -> tuple[list[list[Any]], str | None]:
    suffix = Path(file_name).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        return [list(row) for row in sheet.iter_rows(values_only=True)], sheet.title
    if suffix in {".csv", ".tsv"}:
        text = content.decode("utf-8-sig")
        delimiter = "\t" if suffix == ".tsv" else ","
        return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)], None
    raise GirderImportError("仅支持 .xlsx、.xlsm、.csv 或 .tsv 架梁工点文件。")


def _find_headers(rows: list[list[Any]]) -> tuple[int, list[str]]:
    best_index = 0
    best_score = -1
    for index, row in enumerate(rows[:20]):
        normalized = [_normalize_text(value) for value in row]
        score = sum(any(_normalize_text(alias) == cell for aliases in _HEADER_ALIASES.values() for alias in aliases) for cell in normalized)
        if score > best_score:
            best_index = index
            best_score = score
    return best_index, [str(value or "").strip() for value in rows[best_index]]


def _field_indexes(headers: list[str]) -> dict[str, int]:
    normalized = [_normalize_text(item) for item in headers]
    found: dict[str, int] = {}
    for field, aliases in _HEADER_ALIASES.items():
        alias_set = {_normalize_text(alias) for alias in aliases}
        index = next((idx for idx, value in enumerate(normalized) if value in alias_set), None)
        if index is not None:
            found[field] = index
    return found


def _cell(row: list[Any], index: int) -> Any:
    return row[index] if index < len(row) else None


def _normalize_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def _workpoint_type(value: Any, name: str) -> str:
    text = _normalize_text(value)
    mapping = {
        "桥梁": "bridge",
        "bridge": "bridge",
        "路基": "roadbed",
        "roadbed": "roadbed",
        "隧道": "tunnel",
        "tunnel": "tunnel",
        "涵洞": "culvert",
        "culvert": "culvert",
        "便道": "access",
        "access": "access",
    }
    if text in mapping:
        return mapping[text]
    if "桥" in name:
        return "bridge"
    return "access"


def _side(value: Any) -> str:
    text = _normalize_text(value)
    if text in {"左", "左幅", "left", "l"}:
        return "left"
    if text in {"右", "右幅", "right", "r"}:
        return "right"
    if text in {"双幅", "左右幅", "both", "全幅"}:
        return "both"
    return "unknown"


def _expand_bridge_sides(bridge, side: str, coarse_mode: bool) -> list[tuple[Any, str]]:
    sections_by_side = {item.side: item for item in bridge.work_sections if item.side in {"left", "right"}}
    if side == "both" and sections_by_side:
        return [(sections_by_side[key], key) for key in ("left", "right") if key in sections_by_side]
    if side in sections_by_side:
        return [(sections_by_side[side], side)]
    if side == "unknown" and len(sections_by_side) == 1:
        only_side, section = next(iter(sections_by_side.items()))
        return [(section, only_side)]
    return [(None, side if side != "both" or coarse_mode else "both")]


def _mileage(value: Any) -> float:
    if value in {None, ""}:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().upper().replace(" ", "")
    match = re.match(r"K?(\d+)\+(\d+(?:\.\d+)?)", text)
    if match:
        return float(match.group(1)) * 1000 + float(match.group(2))
    try:
        return float(text)
    except ValueError as exc:
        raise GirderImportError(f"无法识别里程值：{value}") from exc


def _boolean(value: Any, *, default: bool) -> bool:
    if value in {None, ""}:
        return default
    return _normalize_text(value) in {"1", "true", "yes", "是", "需要", "待架"}


def _date_value(value: Any) -> date | None:
    if value in {None, ""}:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise GirderImportError(f"日期必须使用 YYYY-MM-DD：{value}") from exc


def _linked_refs(value: Any) -> list[PassageConditionRef]:
    refs: list[PassageConditionRef] = []
    for raw in re.split(r"[,，;；\n]+", str(value or "")):
        token = raw.strip()
        if not token:
            continue
        prefix, separator, entity_id = token.partition(":")
        ref_type = prefix if separator and prefix in {"structure", "upper_structure", "milestone"} else "structure"
        refs.append(PassageConditionRef(ref_type=ref_type, entity_id=entity_id if separator else token))
    return refs


def _append_evidence(
    evidence: list[SourceEvidence],
    *,
    file_name: str,
    sheet_name: str | None,
    row_no: int,
    field_path: str,
    original_value: Any,
    normalized_value: Any,
    authority_domain: str,
) -> str:
    evidence_id = stable_id("evidence", {"file": file_name, "sheet": sheet_name, "row": row_no, "field": field_path})
    evidence.append(
        SourceEvidence(
            evidence_id=evidence_id,
            source_type="girder_import",
            authority_domain=authority_domain,
            file_name=file_name,
            sheet_name=sheet_name,
            row_or_region=str(row_no),
            field_path=field_path,
            original_value=original_value,
            normalized_value=normalized_value,
        )
    )
    return evidence_id


def _message(level: str, code: str, message: str, *refs: str) -> ValidationMessage:
    return ValidationMessage(level=level, code=code, message=message, subject_id=refs[0] if refs else None, entity_refs=list(refs))
