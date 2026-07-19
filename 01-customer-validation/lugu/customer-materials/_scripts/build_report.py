from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CUSTOMER_MATERIALS = Path(__file__).resolve().parents[1]
OUTPUT = CUSTOMER_MATERIALS / "泸古项目7月15日上午调研验证记录_20260715.docx"

COLORS = {
    "navy": "1F4D78",
    "blue": "2E74B5",
    "ink": "202939",
    "muted": "667085",
    "light": "F2F4F7",
    "blue_light": "EAF2F8",
    "callout": "F4F6F9",
    "border": "D0D5DD",
    "green": "E8F5E9",
    "green_text": "166534",
    "yellow": "FFF4CC",
    "yellow_text": "7A5A00",
    "red": "FDECEC",
    "red_text": "9B1C1C",
    "white": "FFFFFF",
}

LATIN_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"


def set_run_font(run, size=None, bold=None, color=None, italic=None, latin=LATIN_FONT, cjk=CJK_FONT):
    run.font.name = latin
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color="202939", bold=False):
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D5DD", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    assert sum(widths_dxa) == 9360, sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def style_table(table, widths_dxa, compact=False):
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                set_cell_shading(cell, COLORS["light"])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2 if row_idx else 0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.keep_together = True
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_idx == 0 or col_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=8.6 if compact else 9.2,
                        bold=True if row_idx == 0 else False,
                        color=COLORS["navy"] if row_idx == 0 else COLORS["ink"],
                    )


def add_table(doc, headers, rows, widths_dxa, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, widths_dxa, compact=compact)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.line_spacing = 1
    return table


def add_paragraph_border(paragraph, left=None, bottom=None, fill=None):
    p_pr = paragraph._p.get_or_add_pPr()
    if fill:
        shd = p_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            p_pr.append(shd)
        shd.set(qn("w:fill"), fill)
        shd.set(qn("w:val"), "clear")
    if left or bottom:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        if left:
            node = OxmlElement("w:left")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), str(left.get("size", 18)))
            node.set(qn("w:space"), str(left.get("space", 8)))
            node.set(qn("w:color"), left.get("color", COLORS["blue"]))
            borders.append(node)
        if bottom:
            node = OxmlElement("w:bottom")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), str(bottom.get("size", 12)))
            node.set(qn("w:space"), str(bottom.get("space", 8)))
            node.set(qn("w:color"), bottom.get("color", COLORS["blue"]))
            borders.append(node)


def add_callout(doc, label, text, tone="blue"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = True
    fill = COLORS["callout"] if tone == "blue" else COLORS["yellow"]
    border = COLORS["blue"] if tone == "blue" else COLORS["yellow_text"]
    add_paragraph_border(p, left={"color": border, "size": 22, "space": 10}, fill=fill)
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, size=11, bold=True, color=COLORS["navy"] if tone == "blue" else COLORS["yellow_text"])
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=COLORS["ink"])
    return p


def add_numbering(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    n_id = OxmlElement("w:numId")
    n_id.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(n_id)
    p_pr.append(num_pr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=11, bold=True, color=COLORS["ink"])
        second = p.add_run(text[len(bold_prefix):])
        set_run_font(second, size=11, color=COLORS["ink"])
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=COLORS["ink"])
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True
    return p


def add_body(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True, color=COLORS["ink"])
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11, color=COLORS["ink"], italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=COLORS["ink"], italic=italic)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=COLORS["muted"])
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    end = paragraph.add_run(" 页")
    set_run_font(end, size=8.5, color=COLORS["muted"])


doc = Document()
doc.core_properties.title = "泸古项目上午调研验证记录"
doc.core_properties.subject = "基建智能计划管控产品客户验证"
doc.core_properties.author = "广联达斑马产品团队"
doc.core_properties.keywords = "泸古项目, 客户调研, 计划管控, 产品验证"

section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
set_style_font(normal, 11, COLORS["ink"])
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.widow_control = True

title_style = styles["Title"]
set_style_font(title_style, 23, "000000", True)
title_style.paragraph_format.space_before = Pt(0)
title_style.paragraph_format.space_after = Pt(4)

subtitle_style = styles["Subtitle"]
set_style_font(subtitle_style, 14, "373737", False)
subtitle_style.paragraph_format.space_before = Pt(0)
subtitle_style.paragraph_format.space_after = Pt(16)

for name, size, color, before, after in (
    ("Heading 1", 16, COLORS["blue"], 16, 8),
    ("Heading 2", 13, COLORS["blue"], 12, 6),
    ("Heading 3", 12, COLORS["navy"], 8, 4),
):
    style = styles[name]
    set_style_font(style, size, color, True)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.widow_control = True

header = section.header
hp = header.paragraphs[0]
hp.paragraph_format.space_after = Pt(0)
hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
hr1 = hp.add_run("客户调研验证记录")
set_run_font(hr1, size=8.5, bold=True, color=COLORS["muted"])
hr2 = hp.add_run("\t基建智能计划管控中枢｜内部材料")
set_run_font(hr2, size=8.5, color=COLORS["muted"])

footer = section.footer
fp = footer.paragraphs[0]
fp.paragraph_format.space_before = Pt(0)
fp.paragraph_format.space_after = Pt(0)
add_page_number(fp)

bullet_id = add_numbering(doc, "bullet")
decimal_id = add_numbering(doc, "decimal")

# 首页：memo_masthead
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(12)
title = doc.add_paragraph("泸古项目上午调研验证记录", style="Title")
subtitle = doc.add_paragraph("基建智能计划管控产品方向与联合验证建议", style="Subtitle")

metadata = [
    ("项目", "泸古项目"),
    ("访谈对象", "工程部长 李宗仁"),
    ("调研日期", "2026 年 7 月 15 日上午"),
    ("报告定位", "内部产品验证记录"),
    ("资料状态", "尚未正式收集，可通过微信持续沟通获取"),
    ("验证状态", "方向性验证，尚未形成正式试点和结果复核安排"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, size=10.5, bold=True, color="000000")
    r2 = p.add_run(value)
    set_run_font(r2, size=10.5, color="000000")

rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(6)
rule.paragraph_format.space_after = Pt(10)
add_paragraph_border(rule, bottom={"color": COLORS["blue"], "size": 14, "space": 2})

add_callout(
    doc,
    "产品决策",
    "建议将泸古项目作为联合验证对象，但验证重点应从“重新生成初始计划”调整为“复现既有总控逻辑，并增强月度计划、进度反馈、节点余量预警和滚动调整”。",
)

# 一、调研背景与目标
add_heading(doc, "一、调研背景与目标", 1)
add_body(
    doc,
    "当前基建版仍处于研发和技术路径验证阶段。现有 Demo 已能够验证项目结构建模、任务网络生成、资源排程、里程碑诊断和方案比较，但尚不能直接承接泸古项目完整真实数据，也未打通从计划编制到执行反馈、风险预警和调整发布的价值闭环。",
)
add_body(
    doc,
    "本次调研不是再次泛化收集功能需求，而是通过客户现有计划方法判断：客户真正需要解决的问题是什么、当前 Demo 能验证什么、下一步应选择什么真实场景，以及 7 月 30 日前应形成哪些继续或停止的证据。",
)
add_list_item(doc, "还原客户当前总体计划、资源配置和月度管控方法。", bullet_id)
add_list_item(doc, "判断初始计划、动态管控和风险预警三类产品假设的成立程度。", bullet_id)
add_list_item(doc, "收敛泸古项目真实数据验证范围、资料清单、验收标准和后续行动。", bullet_id)

# 二、总体结论
add_heading(doc, "二、总体结论", 1)
conclusions = [
    "客户并不缺一份初始总体计划。客户已使用 Excel 建立全线总控框架、架梁顺序、关键工点工期、主要工装配置和日期联动公式，粗粒度的自动计划无法形成明显替代价值。",
    "最明确的产品机会在执行阶段。客户希望从总体计划中直接筛选月度形象计划，回收现场实际进度与未完成原因，再判断对后续节点的影响并调整下一周期计划。",
    "产品粒度必须分层。控制性工程需要细化到墩台、关键工序和关键资源；普通工程更适合按桥梁、左右幅或月度完成量管理，过细会脱离现场，过粗则没有管理价值。",
    "风险预警必须围绕节点余量。任务晚于原计划并不必然构成风险，只有当偏差侵蚀架梁、合拢或其他控制节点的可用余量时，才需要形成预警和处置建议。",
    "客户具备继续验证的条件，但尚未形成试点承诺。客户愿意通过微信提供计划、架梁顺序和资源资料，也愿意在效果形成后继续沟通；目前尚未收到资料，也未确定复核责任人和时间。",
]
for item in conclusions:
    add_list_item(doc, item, decimal_id)

add_callout(
    doc,
    "阶段结论",
    "调整后推进。首轮验证不以替代 Excel 或直接上线为目标，而以“真实数据能否复现、月度闭环是否有增量价值、关键节点预警是否符合客户口径”为判断标准。",
    tone="yellow",
)

# 三、客户现行业务流程
add_heading(doc, "三、客户现行业务流程", 1)
add_heading(doc, "3.1 初始总控计划形成方式", 2)
initial_flow = [
    "根据结构物工程量和现场经验工效测算持续时间。",
    "以全线架梁顺序、连续梁/连续刚构合拢及最晚完成时间约束关键工点。",
    "结合模板、钻机、挂篮、梁场生产和存梁能力配置主要资源与周转材料。",
    "通过 Excel 公式连接开始日期、结束日期和持续时间，调整关键输入后联动顺延。",
    "对普通工点保留时间弹性，对控制性工点明确最迟完成时间和必须优先施工的对象。",
]
for item in initial_flow:
    add_list_item(doc, item, decimal_id)

add_heading(doc, "3.2 月度执行和进度分析方式", 2)
monthly_flow = [
    "工程部从总体计划中形成当月形象进度目标，通常以完成多少根桩、多少道系梁、多少米墩柱或某个控制墩台来表达。",
    "现场技术员核查实际形象进度，反馈完成数量、累计完成比例和未完成原因。",
    "未完成内容顺延到下一个月，并结合剩余工期要求施工队采取增加资源或压缩周期等措施。",
    "项目通过月度进度分析判断偏差是否影响总工期、架梁通道或其他控制节点。",
    "产值采用形象进度对应的综合单价口径，用于快速形成计划产值和完成产值。",
]
for item in monthly_flow:
    add_list_item(doc, item, decimal_id)

add_heading(doc, "3.3 访谈中出现的角色与职责", 2)
roles = [
    ["工程部长", "编制和维护总体计划，确定架梁及控制节点逻辑，配置主要工装，下达月度计划并组织进度分析。", "Excel 总控计划、资源配置表、架梁顺序资料", "减少重复拆解计划，提高滚动调整和风险判断效率。"],
    ["现场技术员", "核查施工形象进度，反馈完成情况和未完成原因，确认现场实际施工顺序。", "日报、月报、现场进度统计", "降低填报负担，允许实际顺序与初始计划不同。"],
    ["计划/商务相关人员", "依据工程部提供的完成量形成产值和计量相关口径。", "形象进度量、综合单价、产值报表", "保持计划量、实际量和产值口径可对应。"],
]
add_table(doc, ["角色", "主要职责", "当前资料/工具", "可能接受的产品价值"], roles, [1512, 2952, 2232, 2664])

# 四、客户关键表达归纳
add_heading(doc, "四、客户关键表达归纳", 1)
expressions = [
    ("现有计划并非空白", "客户已经建立项目整体计划框架，能够表达架梁顺序、桥梁左右幅关系、关键工点持续时间、工装配置及日期联动。产品若只输出一份粗略横道计划，客户认为 Excel 已足够。"),
    ("计划粒度需要“重点细、普通粗”", "控制性工程应细化到墩台和关键工序；普通桩基不必固定到单根施工次序，可按一个作业单元或月度完成量管理。现场只要在时间窗口内完成，并不要求完全按照系统给出的对象顺序施工。"),
    ("架梁是全线总控的重要约束", "架梁路线需要考虑左右幅穿插、架桥机横移或拆装、隧道通行、连续结构合拢、桥面后续施工和梁场产存能力。部分工点可以延后，前提是不影响架梁的最晚需要时间。"),
    ("风险不是普通日期偏差", "工点晚开但仍有可压缩余量、且不影响架梁时，不应频繁预警；只有偏差突破最晚完成时间或侵蚀控制节点余量时，才需要升级处理。"),
    ("资源必须综合考虑", "资源配置不能只增加钻机或模板数量，还要同时考虑钢筋加工、混凝土供应、班组能力、设备故障、天气、节假日和材料供应等条件，否则会出现等待或窝工。"),
    ("真正希望得到月度闭环", "客户希望总体计划确定后，系统能够筛选当月计划、按合适粒度下达给施工队伍，回收实际进度后锁定已完成事实，再对剩余任务进行调整。"),
    ("使用意愿取决于是否减少额外工作", "客户不希望为了验证软件重新录入大量已有基础数据，也不希望软件开发占用其日常工作时间。可接受的方式是复用既有平台数据，通过微信补充现有计划资料，在形成效果后集中复核。"),
]
for title_text, detail in expressions:
    add_heading(doc, title_text, 3)
    add_body(doc, detail)

# 五、问题与产品机会判断
add_heading(doc, "五、问题与产品机会判断", 1)
add_heading(doc, "5.1 显性问题", 2)
for item in [
    "总体计划已经形成，但每月仍需人工筛选、转换成可下达的形象进度计划。",
    "实际施工顺序与初始计划可能不同，需要在保留已完成事实的基础上重新安排剩余工作。",
    "偏差分析不能只比较计划日期，需要判断是否影响架梁、合拢等控制节点。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "5.2 隐性问题", 2)
for item in [
    "关键计划逻辑高度依赖工程部长个人经验，架梁穿插、连续结构同步和资源流转不容易标准化复用。",
    "Excel 在单次编制时足够灵活，但多轮实际进度更新、影响传播、版本留痕和方案比较成本会持续增加。",
    "计划粒度、现场填报粒度和管理层关注粒度不同，若系统只提供一种颗粒度，容易出现数据无法回收或结果无法使用。",
]:
    add_list_item(doc, item, bullet_id)

add_callout(
    doc,
    "核心矛盾",
    "计划计算需要足够细，现场执行需要保留灵活性。产品不能把所有任务都锁成唯一施工顺序，而应区分必须遵守的硬逻辑、关键节点和资源上限，以及允许在时间窗口内调整的现场执行顺序。",
)

add_heading(doc, "5.3 当前产品机会", 2)
for item in [
    "以客户现有总控计划为基准，而不是要求从零重新生成。",
    "建立架梁总控和单体详细计划的上下层关联，自动计算最晚完成时间和剩余余量。",
    "按控制性工程与普通工程选择不同输出粒度，直接形成月度形象计划。",
    "接收实际进度后锁定已完成事实，对剩余任务进行可解释的滚动预测和方案比较。",
    "只对影响控制节点的偏差形成风险预警，并给出影响范围和处置窗口。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "5.4 非优先机会点", 2)
for item in [
    "仅生成一份粗略总体计划或通用甘特图。",
    "强制细化并固定普通工程中每根桩、每个墩台的唯一施工顺序。",
    "在资料和现场口径尚未验证前，直接承诺自动接口和生产级上线。",
    "首期优先建设微信群机器人等自动采集方式；客户对此未形成明确使用承诺。",
]:
    add_list_item(doc, item, bullet_id)

# 六、产品假设验证
doc.add_page_break()
add_heading(doc, "六、产品假设验证", 1)
hypotheses = [
    ["初始计划自动生成是首要价值", "弱成立", "客户已具备较成熟的 Excel 总控计划；粗略计划没有增量价值。", "保留初始排程能力，但验证重点转向计划复现、校准和执行闭环。"],
    ["真实项目具备可用的数据基础", "条件成立", "项目已有参数化结构数据、总体计划、资源配置和架梁顺序资料，但尚未正式取得。", "先通过微信获取数据并做质量检查，不能把“资料存在”视为“数据可直接入模”。"],
    ["月度计划可由总体计划自动生成", "强成立", "客户明确希望按月筛选桥梁、墩台或完成量，并直接下达给现场。", "优先验证粗细粒度切换、计划筛选、导出和形象进度口径。"],
    ["实际进度反馈后需要滚动重排", "强成立", "当前流程会把未完成任务顺延到下月，并在进度分析中判断是否采取压缩措施。", "重排必须锁定实际完成事实、允许实际顺序变化，并解释后续节点影响。"],
    ["风险预警能够形成管理价值", "条件成立", "客户只关注突破余量并影响架梁、合拢等控制节点的风险。", "预警应基于最晚完成时间、剩余缓冲和影响传播，避免普通日期偏差产生噪声。"],
    ["客户愿意进入真实项目试点", "条件成立", "客户愿意通过微信提供资料并在效果形成后继续沟通。", "尚未取得数据、责任人和复核日期，当前不能表述为正式试点已确认。"],
]
hyp_table = add_table(doc, ["产品假设", "验证结果", "判断依据", "产品启示与风险"], hypotheses, [2232, 1152, 3024, 2952], compact=True)
for row in hyp_table.rows[1:]:
    status = row.cells[1].text
    if status == "强成立":
        set_cell_shading(row.cells[1], COLORS["green"])
        color = COLORS["green_text"]
    elif status == "条件成立":
        set_cell_shading(row.cells[1], COLORS["yellow"])
        color = COLORS["yellow_text"]
    else:
        set_cell_shading(row.cells[1], COLORS["light"])
        color = COLORS["muted"]
    for run in row.cells[1].paragraphs[0].runs:
        set_run_font(run, size=8.6, bold=True, color=color)

add_callout(
    doc,
    "验证重点",
    "当前最强证据支持月度计划生成和实际进度滚动管控；最大不确定性仍是客户数据能否快速入模，以及客户是否愿意安排正式结果复核。",
)

# 七、当前 Demo 匹配与缺口
add_heading(doc, "七、当前 Demo 匹配与缺口", 1)
add_body(
    doc,
    "当前 Demo 已验证项目结构、工艺工效、施工逻辑、资源和里程碑的统一建模，并具备任务网络生成、固定资源求工期、目标工期测算资源、方案比较和主要瓶颈诊断能力。这些能力可以作为泸古项目验证的计算基础。",
)
demo_rows = [
    ["P0", "客户既有计划接入", "当前以系统生成计划为主，客户已有成熟 Excel 基准。", "支持将既有计划作为基准版本，保留客户确定的关键时间和调整记录。"],
    ["P0", "架梁总控逻辑", "现有简单架梁假设不能完整表达左右幅穿插、往返、横移/拆装、穿隧道和连续结构制约。", "先复现客户确认的架梁路线，再计算各工点最晚完成时间。"],
    ["P0", "粗细粒度切换", "当前任务网络偏细，普通桩基按单根排序容易与现场脱节。", "控制工程细到墩台/关键工序，普通工程按桥、幅、作业单元或完成量输出。"],
    ["P0", "月度计划与实际反馈", "月度计划输出、实际进度锁定和剩余任务重排尚未形成稳定闭环。", "以一个月度周期做 MVP，先支持人工或 Excel 录入。"],
    ["P1", "关键节点风险预警", "现有偏差诊断尚未完全体现最晚完成时间和剩余缓冲。", "只对侵蚀架梁/合拢余量的偏差预警，并展示影响链和处置窗口。"],
    ["P1", "工艺和资源模板适配", "柱系梁与墩柱、连续梁/连续刚构同步等项目逻辑仍需校准。", "把项目特有逻辑作为可追溯规则，不写死为通用行业规则。"],
    ["P2", "自动进度采集", "微信群等自动采集方式存在内部数据和使用意愿限制。", "首期使用人工/Excel，待验证填报成本后再判断自动采集。"],
]
demo_table = add_table(doc, ["优先级", "能力", "当前缺口", "建议处理"], demo_rows, [792, 2448, 3024, 3096], compact=True)
for row in demo_table.rows[1:]:
    priority = row.cells[0].text
    if priority == "P0":
        set_cell_shading(row.cells[0], COLORS["red"])
        color = COLORS["red_text"]
    elif priority == "P1":
        set_cell_shading(row.cells[0], COLORS["yellow"])
        color = COLORS["yellow_text"]
    else:
        set_cell_shading(row.cells[0], COLORS["light"])
        color = COLORS["muted"]
    for run in row.cells[0].paragraphs[0].runs:
        set_run_font(run, size=8.6, bold=True, color=color)

# 八、下一步真实项目验证方案
doc.add_page_break()
add_heading(doc, "八、下一步真实项目验证方案", 1)
add_heading(doc, "8.1 验证范围", 2)
add_callout(
    doc,
    "推荐场景",
    "全线架梁总控 + 一座控制性大桥详细计划 + 一个完整月度计划周期。",
)
add_body(
    doc,
    "该范围同时保留客户最关注的全线控制逻辑和单体排程验证，又能通过一个月度周期验证实际进度反馈与风险判断，避免一次覆盖全项目所有专业。",
)

add_heading(doc, "8.2 微信资料收集清单", 2)
for item in [
    "当前有效的总体进度计划 Excel，并确认计划版本和状态日期。",
    "全线架梁顺序图或对应表格，包括左右幅顺序、架桥机横移/拆装、隧道通行和控制节点。",
    "所选控制性大桥的结构、工程量、关键工效和工艺逻辑。",
    "主要资源配置表，包括钻机、模板、挂篮、梁场生产和存梁能力。",
    "最近一期月度计划、实际完成情况和未完成原因样例。",
    "现有参数化模型或平台数据的可用范围及导出方式。",
]:
    add_list_item(doc, item, bullet_id)
add_body(doc, "资料获取状态：以上资料尚未正式收到，建议通过微信发送最小数据清单并逐项确认；所有日期均为内部建议，需要客户确认。", bold_prefix="资料获取状态：")

add_heading(doc, "8.3 验证步骤", 2)
validation_steps = [
    "资料接收与映射：记录数据来源、缺失项、冲突项和人工假设。",
    "计划复现：先复现客户架梁路线、控制节点和资源配置，不急于生成“更优方案”。",
    "结果校准：由客户复核关键架梁逻辑、控制桥持续时间和最晚完成时间。",
    "月度闭环：从基准计划生成一个月度计划，输入一次实际偏差，锁定实际后重排剩余任务。",
    "价值判断：比较现有 Excel 与系统在计划筛选、影响分析、方案调整和解释方面的增量价值。",
]
for item in validation_steps:
    add_list_item(doc, item, decimal_id)

add_heading(doc, "8.4 验收标准", 2)
criteria = [
    ["架梁总控", "客户确认的架梁路线、左右幅穿插、横移/拆装及关键制约能够正确表达。"],
    ["节点可信度", "控制节点和最晚完成时间与客户计划差异可解释，不能以隐藏默认值掩盖分歧。"],
    ["计划粒度", "控制性工程能够细化，普通工程能够按桥、幅、作业单元或完成量输出。"],
    ["月度闭环", "能够形成一个月度计划，录入实际进度后锁定事实并更新剩余计划。"],
    ["预警质量", "普通偏差不误报；当偏差影响架梁或合拢余量时，能够说明影响对象和处置窗口。"],
    ["继续意愿", "客户完成一次结果复核，并明确继续校准、进入试点或停止当前路径。"],
]
add_table(doc, ["验证维度", "通过标准"], criteria, [1800, 7560])

add_heading(doc, "8.5 建议行动项", 2)
actions = [
    ["A01", "通过微信发送最小数据清单，确认资料版本、联系人和可提供时间。", "我方产品负责人", "7 月 16 日", "未启动"],
    ["A02", "提供总体计划、架梁顺序、资源配置及一个月度计划样例。", "李宗仁/客户", "建议 7 月 18 日", "待客户确认"],
    ["A03", "完成数据检查、映射和问题清单，确定控制性大桥范围。", "我方产品/算法", "7 月 21 日", "未启动"],
    ["A04", "完成首轮计划复现、月度计划输出和偏差模拟。", "我方产品/算法", "7 月 24 日", "未启动"],
    ["A05", "组织客户结果复核，记录认可、需校准和不认可项。", "双方", "建议 7 月 28 日", "待客户确认"],
    ["A06", "形成绿/黄/红验证结论和下一阶段产品决策。", "我方产品负责人", "7 月 30 日", "未启动"],
]
action_table = add_table(doc, ["编号", "行动项", "责任人", "建议时间", "状态"], actions, [864, 3456, 1800, 1800, 1440], compact=True)
for row in action_table.rows[1:]:
    status = row.cells[4].text
    if "待客户确认" in status:
        set_cell_shading(row.cells[4], COLORS["yellow"])
        for run in row.cells[4].paragraphs[0].runs:
            set_run_font(run, size=8.6, bold=True, color=COLORS["yellow_text"])

# 九、下一轮需确认问题
add_heading(doc, "九、下一轮需确认问题", 1)
questions = [
    "客户当前有效的总体计划版本、状态日期和开工调整口径是什么？",
    "哪一座桥最适合作为控制性大桥样例，其架梁或合拢目标日期是什么？",
    "现有参数化模型、Excel 总控计划和现场月报中，哪一份是各类数据的权威来源？",
    "普通工程月度计划更适合按桥、左右幅、作业单元还是完成量下达？",
    "客户能够接受的日期、工期和资源结果误差范围是什么？哪些差异必须逐项解释？",
    "谁负责复核任务逻辑、资源配置和结果价值，何时可以安排首次结果沟通？",
]
for item in questions:
    add_list_item(doc, item, decimal_id)

# 十、最终建议
add_heading(doc, "十、最终建议", 1)
add_body(
    doc,
    "泸古项目具备较强的计划专业能力和相对完整的数据基础，是验证产品是否真正理解复杂施工组织的合适对象。但客户现有 Excel 已能解决初始总控计划问题，继续以“自动生成计划”为主要卖点，难以形成明显价值。",
)
add_body(
    doc,
    "下一步应采用高介入的联合验证方式：人工完成首轮数据映射，先复现客户现有计划，再验证月度计划、实际进度反馈、节点余量预警和滚动调整。只有客户确认系统结果可信、能够减少计划拆解或影响分析工作，并愿意安排持续复核，才能认为价值闭环开始成立。",
)
add_callout(
    doc,
    "建议",
    "在收到最小数据包前，不继续扩大 Demo 功能范围；收到资料后，以“复现优先、解释优先、一个闭环优先”为原则完成首轮真实验证。",
    tone="yellow",
)
add_body(
    doc,
    "本报告结论仅基于本次上午访谈材料。客户真实数据质量、系统计算结果、实际使用频率、组织投入和正式试点意愿仍需通过下一轮验证确认。",
    italic=True,
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
