from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ERROR_VALUES = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
BRIDGE_TERMS = ("桥", "互通", "隧道", "路基", "梁场")
CONTROL_TERMS = ("关键", "节点", "架梁", "合拢", "完工", "贯通", "梁场")
RESOURCE_TERMS = ("钻机", "模板", "挂篮", "台座", "设备", "班组", "工装", "产能", "资源")
COMPONENT_RE = re.compile(r"(?:\d+\s*[#号]\s*(?:墩|台|桩)|(?:墩|台|桩)\s*\d+\s*[#号]?)")


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def meaningful_text(values: list[Any]) -> str:
    return " | ".join(str(value).strip() for value in values if value not in (None, ""))


def inspect_workbook(path: Path, include_samples: bool = False) -> dict[str, Any]:
    wb = load_workbook(path, read_only=False, data_only=False)
    wb_cached = load_workbook(path, read_only=False, data_only=True)

    totals = Counter()
    sheets: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        cached_ws = wb_cached[ws.title]
        stats = Counter()
        component_rows = 0
        bridge_rows = 0
        control_rows = 0
        resource_rows = 0
        samples: list[list[Any]] = []
        formula_errors: list[dict[str, Any]] = []

        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            values = [cell.value for cell in row]
            nonempty = [value for value in values if value not in (None, "")]
            if nonempty:
                stats["nonempty_rows"] += 1
                text = meaningful_text(values)
                if COMPONENT_RE.search(text):
                    component_rows += 1
                if any(term in text for term in BRIDGE_TERMS):
                    bridge_rows += 1
                if any(term in text for term in CONTROL_TERMS):
                    control_rows += 1
                if any(term in text for term in RESOURCE_TERMS):
                    resource_rows += 1
                if include_samples and len(samples) < 90:
                    samples.append([json_value(value) for value in values])

            for col_idx, cell in enumerate(row, start=1):
                value = cell.value
                if value not in (None, ""):
                    stats["nonempty_cells"] += 1
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    stats["formula_cells"] += 1
                    cached_value = cached_ws.cell(row=row_idx, column=col_idx).value
                    if cached_value in ERROR_VALUES:
                        formula_errors.append(
                            {
                                "cell": cell.coordinate,
                                "formula": value,
                                "cached": cached_value,
                            }
                        )
                if isinstance(value, (datetime, date)):
                    stats["date_cells"] += 1
                if isinstance(value, str) and value in ERROR_VALUES:
                    stats["literal_errors"] += 1

        sheet = {
            "name": ws.title,
            "state": ws.sheet_state,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": len(ws.merged_cells.ranges),
            "nonempty_rows": stats["nonempty_rows"],
            "nonempty_cells": stats["nonempty_cells"],
            "formula_cells": stats["formula_cells"],
            "date_cells": stats["date_cells"],
            "formula_error_count": len(formula_errors),
            "literal_error_count": stats["literal_errors"],
            "component_reference_rows": component_rows,
            "bridge_or_workpoint_rows": bridge_rows,
            "control_term_rows": control_rows,
            "resource_term_rows": resource_rows,
        }
        if formula_errors:
            sheet["formula_errors"] = formula_errors[:30]
        if include_samples:
            sheet["sample_nonempty_rows"] = samples
        sheets.append(sheet)
        totals.update(
            {
                "sheets": 1,
                "rows": ws.max_row,
                "columns_sum": ws.max_column,
                "nonempty_rows": stats["nonempty_rows"],
                "nonempty_cells": stats["nonempty_cells"],
                "formula_cells": stats["formula_cells"],
                "date_cells": stats["date_cells"],
                "formula_error_count": len(formula_errors),
                "literal_error_count": stats["literal_errors"],
                "component_reference_rows": component_rows,
                "bridge_or_workpoint_rows": bridge_rows,
                "control_term_rows": control_rows,
                "resource_term_rows": resource_rows,
            }
        )

    wb.close()
    wb_cached.close()
    return {
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "sheet_names": [sheet["name"] for sheet in sheets],
        "totals": dict(totals),
        "sheets": sheets,
    }


def reduction(old: int, new: int) -> float | None:
    if old == 0:
        return None
    return round((old - new) / old * 100, 1)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--old", required=True, type=Path)
    parser.add_argument("--new", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    old = inspect_workbook(args.old, include_samples=False)
    new = inspect_workbook(args.new, include_samples=True)
    comparison = {
        "size_reduction_pct": reduction(old["size_bytes"], new["size_bytes"]),
        "sheet_reduction_pct": reduction(old["totals"]["sheets"], new["totals"]["sheets"]),
        "row_reduction_pct": reduction(old["totals"]["rows"], new["totals"]["rows"]),
        "nonempty_cell_reduction_pct": reduction(
            old["totals"]["nonempty_cells"], new["totals"]["nonempty_cells"]
        ),
        "formula_reduction_pct": reduction(
            old["totals"]["formula_cells"], new["totals"]["formula_cells"]
        ),
    }
    payload = {"old": old, "new": new, "comparison": comparison}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"comparison": comparison, "old": old["totals"], "new": new["totals"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
