from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--sheet-index", type=int, default=2)
    parser.add_argument("--start", type=int, default=1)
    parser.add_argument("--end", type=int, default=9999)
    args = parser.parse_args()

    wb = load_workbook(args.input, read_only=True, data_only=False)
    ws = wb.worksheets[args.sheet_index]
    print(f"SHEET={ws.title}")
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if idx < args.start or idx > args.end:
            continue
        values = [str(value).replace("\n", " / ") for value in row[:9] if value not in (None, "")]
        if values:
            print(f"{idx}: " + " | ".join(values))
    wb.close()


if __name__ == "__main__":
    main()
