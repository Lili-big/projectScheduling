from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "泸古项目计划粒度验证记录_20260716_v2.docx"

COLORS = {
    "navy": "1F4D78",
    "blue": "2E74B5",
    "ink": "202939",
    "muted": "667085",
    "light": "F2F4F7",
    "callout": "F4F6F9",
    "border": "D0D5DD",
    "green": "E8F5E9",
    "green_text": "166534",
    "yellow": "FFF4CC",
    "yellow_text": "7A5A00",
    "red": "FDECEC",
    "red_text": "9B1C1C",
}
LATIN_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color="202939", bold=False):
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)


def add_paragraph_border(paragraph, left=None, bottom=None, fill=None):
    p_pr = paragraph._p.get_or_add_pPr()
    if fill:
        shd = p_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            p_pr.append(shd)
        shd.set(qn("w:fill"), fill)
        shd.set(qn("w:val"), "clear")
    if left or bottom:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        if left:
            node = OxmlElement("w:left")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), str(left.get("size", 18)))
            node.set(qn("w:space"), str(left.get("space", 8)))
            node.set(qn("w:color"), left.get("color", COLORS["blue"]))
            borders.append(node)
        if bottom:
            node = OxmlElement("w:bottom")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), str(bottom.get("size", 12)))
            node.set(qn("w:space"), str(bottom.get("space", 8)))
            node.set(qn("w:color"), bottom.get("color", COLORS["blue"]))
            borders.append(node)


def add_callout(doc, label, text, tone="blue"):
    palette = {
        "blue": (COLORS["callout"], COLORS["blue"], COLORS["navy"]),
        "yellow": (COLORS["yellow"], COLORS["yellow_text"], COLORS["yellow_text"]),
        "green": (COLORS["green"], COLORS["green_text"], COLORS["green_text"]),
        "red": (COLORS["red"], COLORS["red_text"], COLORS["red_text"]),
    }
    fill, border, label_color = palette[tone]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = True
    add_paragraph_border(p, left={"color": border, "size": 22, "space": 10}, fill=fill)
    set_run_font(p.add_run(f"{label}："), size=11, bold=True, color=label_color)
    set_run_font(p.add_run(text), size=11, color=COLORS["ink"])
    return p


def add_numbering(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    n_id = OxmlElement("w:numId")
    n_id.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(n_id)
    p_pr.append(num_pr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=11, bold=True, color=COLORS["ink"])
        set_run_font(p.add_run(text[len(bold_prefix):]), size=11, color=COLORS["ink"])
    else:
        set_run_font(p.add_run(text), size=11, color=COLORS["ink"])
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True
    return p


def add_body(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=11, bold=True, color=COLORS["ink"])
        set_run_font(p.add_run(text[len(bold_prefix):]), size=11, color=COLORS["ink"], italic=italic)
    else:
        set_run_font(p.add_run(text), size=11, color=COLORS["ink"], italic=italic)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    assert sum(widths_dxa) == 9360
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), COLORS["border"])
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths_dxa, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    set_table_geometry(table, widths_dxa)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, COLORS["light"])
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2 if r_idx else 0)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_together = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or c_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(
                        run,
                        size=8.5 if compact else 9.0,
                        bold=(r_idx == 0),
                        color=COLORS["navy"] if r_idx == 0 else COLORS["ink"],
                    )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(3)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraph.add_run("第 "), size=8.5, color=COLORS["muted"])
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(paragraph.add_run(" 页"), size=8.5, color=COLORS["muted"])


doc = Document()
doc.core_properties.title = "泸古项目计划粒度验证记录"
doc.core_properties.subject = "基于既有桥梁排程文件与客户反馈的产品验证"
doc.core_properties.author = "广联达斑马产品团队"
doc.core_properties.keywords = "泸古项目, 计划粒度, 客户验证, 桥梁排程, 形象进度"

section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
set_style_font(normal, 11, COLORS["ink"])
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.widow_control = True
title_style = styles["Title"]
set_style_font(title_style, 23, "000000", True)
title_style.paragraph_format.space_before = Pt(0)
title_style.paragraph_format.space_after = Pt(4)
subtitle_style = styles["Subtitle"]
set_style_font(subtitle_style, 14, "373737", False)
subtitle_style.paragraph_format.space_before = Pt(0)
subtitle_style.paragraph_format.space_after = Pt(16)
for name, size, color, before, after in (
    ("Heading 1", 16, COLORS["blue"], 16, 8),
    ("Heading 2", 13, COLORS["blue"], 12, 6),
    ("Heading 3", 12, COLORS["navy"], 8, 4),
):
    style = styles[name]
    set_style_font(style, size, color, True)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.widow_control = True

header = section.header
hp = header.paragraphs[0]
hp.paragraph_format.space_after = Pt(0)
hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
set_run_font(hp.add_run("客户验证记录"), size=8.5, bold=True, color=COLORS["muted"])
set_run_font(hp.add_run("\t泸古项目｜内部材料"), size=8.5, color=COLORS["muted"])
footer = section.footer
fp = footer.paragraphs[0]
fp.paragraph_format.space_before = Pt(0)
fp.paragraph_format.space_after = Pt(0)
add_page_number(fp)

bullet_id = add_numbering(doc, "bullet")
decimal_id = add_numbering(doc, "decimal")

# 首页：standard_business_brief + memo_masthead
doc.add_paragraph().paragraph_format.space_after = Pt(12)
doc.add_paragraph("泸古项目计划粒度验证记录", style="Title")
doc.add_paragraph("总控计划收缩行为与分层计划方向跟进验证｜V2", style="Subtitle")
metadata = [
    ("项目", "泸古高速 TJ-1 标"),
    ("客户角色", "工程部长"),
    ("验证日期", "2026 年 7 月 16 日"),
    ("新增材料", "《泸古高速项目TJ-1标总体进度计划-1标6.6.xlsx》"),
    ("验证主题", "客户主动收缩计划粒度是否支持分层计划方向"),
    ("验证状态", "方向证据增强，进入最小场景验证"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run(f"{label}："), size=10.5, bold=True, color="000000")
    set_run_font(p.add_run(value), size=10.5, color="000000")
rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(6)
rule.paragraph_format.space_after = Pt(10)
add_paragraph_border(rule, bottom={"color": COLORS["blue"], "size": 14, "space": 2})
add_callout(
    doc,
    "验证结论",
    "建议按分层计划方向继续推进。新版计划已经从15张逐桥明细表主动收缩为4张总控及关键工点表：普通工程按工点、左右幅和阶段控制，永宁河特大桥主桥、两河口大桥等控制性工程仍细化到主墩和关键工序。这一变化用实际编表行为确认了“项目总控＋控制性工程精排＋普通工程形象进度”的方向。",
    tone="green",
)

add_heading(doc, "一、验证背景与目标", 1)
add_body(doc, "客户此前尝试将桥梁工点逐个拆解，按照桩基、承台、系梁、墩柱、盖梁等对象逐项计算工期、资源和开始结束时间。该方式能够形成完整明细排程，但工程部长反馈其工作量和调整成本过高，已不再按该方式维护全线计划。")
add_body(doc, "本次新增取得客户6月6日版总体进度计划。验证重点由“客户是否排斥精细计划”进一步收敛为：客户在实际重编计划时保留了哪些计算逻辑、主动删除了哪些维护对象，以及产品应如何承接客户现有总控计划并按需展开控制性工程。")

add_heading(doc, "二、验证材料与方法", 1)
add_heading(doc, "2.1 验证材料", 2)
for item in [
    "《泸古高速TJ-1标桥梁进度统计表4.27(1).xlsx》：客户此前形成的桥梁精细排程文件。",
    "《泸古高速项目TJ-1标总体进度计划-1标6.6.xlsx》：客户收缩后的全标段总控及关键工点计划。",
    "工程部长反馈：明细方式过于复杂、工作量大、维护和调整成本过高。",
    "工程部长反馈：除控制性工程外，普通桥墩主要按形象进度量管理，不会严格控制到具体墩台。",
    "上午调研记录：客户对架梁总控、控制节点、月度计划和现场执行粒度的补充说明。",
]:
    add_list_item(doc, item, bullet_id)
add_heading(doc, "2.2 验证方法", 2)
for item in [
    "对比两版工作簿的工作表数量、数据规模、公式数量和任务拆分层级。",
    "对照客户反馈，判断表格粒度与现场实际管理粒度是否一致。",
    "识别新版保留的工期、资源和施工逻辑，区分“计算价值”与“用户维护粒度”。",
    "形成后续产品能力收敛和真实客户验证建议。",
]:
    add_list_item(doc, item, decimal_id)

add_heading(doc, "三、新旧计划版本对比", 1)
add_body(doc, "客户并未放弃工期策划，而是把计划从“逐桥明细计算表”重构为“全标段总控＋各专业及关键工点”。以下数据为程序读取两版工作簿所得。")
add_table(
    doc,
    ["对比项", "4.27 逐桥明细版", "6.6 总控计划版", "变化"],
    [
        ["工作表", "15 张", "4 张", "减少 73.3%"],
        ["文件大小", "329,671 字节", "71,914 字节", "减少 78.2%"],
        ["累计行数", "1,130 行", "613 行", "减少 45.8%"],
        ["非空单元格", "16,499 个", "3,829 个", "减少 76.8%"],
        ["公式", "5,030 个", "1,450 个", "减少 71.2%"],
        ["具体墩台/桩号引用行", "约 295 行", "约 22 行", "主要集中在控制性工程"],
        ["公式缓存错误", "13 个", "未检出", "仅说明保存时未见缓存错误，不代表逻辑已复核"],
    ],
    [1700, 2300, 2300, 3060],
    compact=True,
)
add_heading(doc, "3.1 新版计划保留了什么", 2)
for item in [
    "保留工点持续时间、开始时间、结束时间、是否关键工点和计算依据，说明工期计算仍是策划基础。",
    "保留1#、2#梁场的产能、台座、模板、存梁能力和全线架梁顺序，说明资源与施工路线仍需统一总控。",
    "保留路基、桥梁、隧道、互通等专业计划和标段总体完工节点，形成跨专业的总控视图。",
    "对永宁河特大桥主桥、两河口大桥等控制性工程，仍展开到主墩桩基、承台、墩身、0号块、挂篮、悬浇段和合龙。",
]:
    add_list_item(doc, item, bullet_id)
add_heading(doc, "3.2 新版计划主动收缩了什么", 2)
for item in [
    "普通桥梁不再逐桩逐墩列任务，通常收缩为左/右幅下部结构、架梁和桥面系等3至6项。",
    "普通工程的构件数量、模板和设备更多作为工期计算依据或备注，不再全部变成需要持续维护的日期任务。",
    "全线计划不再按一桥一表分散维护，而是集中到总体工点、专业和关键节点层级。",
    "控制性工程与普通工程采用不同粒度，说明客户已经实际执行了分层计划，而不是仅表达偏好。",
]:
    add_list_item(doc, item, bullet_id)
add_callout(doc, "新增证据判断", "新版文件是比口头反馈更强的行为证据：客户需要的是可计算的总控计划，但不接受把全部计算中间对象都转化为长期维护任务。产品应让计算保持精细、管理界面按需聚合。", tone="green")

add_heading(doc, "四、客户反馈与行为证据归纳", 1)
add_heading(doc, "4.1 客户明确表达", 2)
for item in [
    "初次编制工作量大，需要逐桥、逐墩、逐工序录入工程量、工效、资源和日期。",
    "后续维护困难，一项条件变化会引起大量任务和日期的连锁调整。",
    "普通桥梁现场不会严格按照具体墩台顺序施工，过细计划与现场执行脱节。",
    "非关键工程更适合控制月度完成多少根桩、多少道系梁、多少米墩柱等形象进度量。",
]:
    add_list_item(doc, item, bullet_id)
add_heading(doc, "4.2 隐含管理原则", 2)
for item in [
    "计划粒度必须与管理决策粒度一致，不是越细越好。",
    "控制性工程需要精确到墩台和关键工序，因为其偏差会影响架梁、合拢或总工期。",
    "普通工程只需要明确时间窗口、完成量和是否影响后续节点，具体顺序应允许现场调整。",
    "计划调整的价值在于快速判断影响和提出措施，而不是重新维护全部明细日期。",
]:
    add_list_item(doc, item, bullet_id)
add_heading(doc, "4.3 客户实际编表行为", 2)
for item in [
    "将关键工点作为显式字段，在27个总控工点中标记3个关键工点。",
    "对普通桥梁采用桥/幅/阶段计划，对连续刚构等控制性工程采用主墩/关键工序计划。",
    "继续用工程量、工效、设备和施工逻辑计算工期，但大幅减少用户必须维护的任务行和跨表公式。",
    "将架梁顺序、梁场产能和控制节点放在全线视角统一管理，符合工程部长上午表达的总控思路。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "五、核心问题判断", 1)
add_callout(doc, "核心矛盾", "原方案把工期计算所需的工程明细，直接等同于用户每天要维护的计划任务。新版证明客户并未放弃计算，而是把计算依据留在后台，把日常控制对象收缩到总控工点、关键节点和少量控制性工程。", tone="yellow")
add_table(
    doc,
    ["维度", "原有方式", "客户实际需要", "验证判断"],
    [
        ["计划范围", "所有桥梁统一精细化", "全标段总控，关键工程精排", "统一粒度不成立"],
        ["任务粒度", "逐桩、逐墩、逐工序", "重点到墩，普通按桥/幅/阶段", "分层方式已在新版出现"],
        ["日期控制", "每项任务固定开始结束日期", "普通任务控制时间窗口", "需保留现场弹性"],
        ["调整方式", "条件变化后全面联动修改", "优先判断控制节点影响", "应按影响范围重排"],
        ["结果用途", "形成完整明细横道计划", "支持节点、架梁总控和月度管理", "结果必须可维护、可执行"],
    ],
    [1500, 2520, 3000, 2340],
    compact=True,
)
add_body(doc, "因此，本轮可以更明确地排除“客户不需要工期计算”的解释。客户新版仍包含大量持续时间公式、工效和资源依据；被否定的是把全部计算中间对象都变成前台计划任务，并要求长期逐项维护。")

add_heading(doc, "六、产品假设验证结果", 1)
add_table(
    doc,
    ["产品假设", "结果", "主要依据", "产品启示"],
    [
        ["全线所有桥梁统一精排到墩台和构件", "明确不成立", "新版主动减少73.3%的工作表和71.2%的公式", "不能作为默认产品模式"],
        ["以总控计划管理项目节点", "强成立", "客户已明确将6.6版作为节点管控计划", "应直接承接为基准计划"],
        ["控制性工程精排到墩台和关键工序", "强成立", "新版对永宁河、两河口主桥保留主墩及合龙明细", "保留精排引擎并按需展开"],
        ["普通桥梁按桥/幅/阶段聚合管理", "强成立", "新版多数普通桥梁收缩为下构、架梁、桥面系", "前台按管理粒度展示"],
        ["同一项目支持粗细粒度切换", "强成立", "新版已同时使用总控、专业、普通桥梁和控制工程粒度", "作为 P0 产品能力"],
        ["全量明细自动重排即可解决动态调整", "不成立", "自动计算不能消除过细结果的维护成本", "重排范围和输出粒度必须可控"],
        ["开工后按月滚动调整", "材料不足，暂不判断", "项目尚未正式开工，缺少实际周期验证", "后续用真实月度数据验证"],
    ],
    [2500, 1500, 2880, 2480],
    compact=True,
)
add_callout(doc, "阶段结论", "分层计划方向已由客户新版计划和明确管理口径共同验证；尚未验证的是产品能否无损承接该计划、完成一次可信调整，以及开工后的月度滚动与节点预警是否产生持续价值。", tone="green")

add_heading(doc, "七、建议的产品方向", 1)
add_heading(doc, "7.1 建立三级计划模型", 2)
add_table(
    doc,
    ["层级", "管理对象", "建议粒度", "主要输出"],
    [
        ["项目总控层", "全标段专业、工点、梁场及关键节点", "按专业、工点、桥和里程碑", "基准计划、架梁路线、节点余量"],
        ["控制性工程层", "连续梁、关键主墩、架梁通道工程", "细化到墩台和关键工序", "工期、资源、瓶颈和调整方案"],
        ["普通工程执行层", "非关键桥梁和普通工点", "按桥/幅/阶段及月度完成量", "月计划、完成量、偏差和剩余任务"],
    ],
    [1700, 2820, 2480, 2360],
    compact=True,
)
add_heading(doc, "7.2 产品能力优先级", 2)
for item in [
    "P0｜导入并复现客户6.6版总控计划，保留工点、关键标记、日期、计算依据和架梁顺序。",
    "P0｜支持控制性工程标记、层级展开与聚合，避免全项目强制统一拆分。",
    "P0｜支持按桥、左右幅、阶段和月度完成量表达普通工程计划。",
    "P0｜预警围绕架梁、合拢和标段完工等节点余量，不因普通任务日期变化频繁报警。",
    "P1｜控制性工程保留任务网络、工期计算、资源排程和多方案比较。",
    "P1｜进度反馈后只重排受影响范围，并把结果重新汇总到专业、工点和项目总控层。",
    "P2｜逐桩、逐构件明细作为计算、追溯或特殊场景能力，不作为默认维护界面。",
]:
    add_list_item(doc, item, bullet_id, bold_prefix=item.split("｜")[0] + "｜")

add_heading(doc, "八、下一步验证方案", 1)
add_body(doc, "下一轮不再验证“能否从零生成一套全量精细计划”，而应以客户已采用的6.6版总控计划为基准，验证产品是否能够承接、解释和调整这套计划。建议采用“全标段总控复现＋一座控制性桥梁精排＋一座普通桥梁聚合计划”的最小闭环。")
for item in [
    "总控复现：导入工点、关键标记、开始结束时间、计算依据、梁场产能和架梁顺序，确保系统结果与客户当前节点口径一致。",
    "控制性桥梁：优先选择永宁河特大桥主桥或两河口大桥，复现主墩、挂篮、悬浇和合龙逻辑。",
    "普通桥梁：选择一座非控制性桥梁，只保留下部结构、左右幅架梁、桥面系及月度完成量。",
    "变化场景：将实际开工条件调整为 9 月 1 日，判断哪些工点、架梁路径和关键节点受到影响，并只重排影响范围。",
    "客户复核：由总工和工程部长共同评价节点差异是否可解释、维护量是否可接受、调整结果是否可采用。",
]:
    add_list_item(doc, item, decimal_id)

add_heading(doc, "8.1 验证指标", 2)
add_table(
    doc,
    ["验证维度", "判断标准", "证据形式"],
    [
        ["基准复现", "关键工点、日期、架梁顺序和资源依据可完整承接", "系统与6.6版差异清单"],
        ["计算可信度", "控制性工程关键工期和节点差异可解释", "总工/工程部长复核记录"],
        ["调整效率", "开工变化后快速形成受影响范围和新节点", "调整步骤、耗时和人工修改量"],
        ["粒度适配", "控制工程可展开，普通工程保持聚合", "新旧任务对象数量对比"],
        ["节点预警", "只在侵蚀架梁、合龙或总工期余量时预警", "场景结果及客户判断"],
        ["持续使用意愿", "客户愿意完成下一轮数据更新和结果复核", "责任人、资料和复核时间"],
    ],
    [1900, 4300, 3160],
    compact=True,
)

add_heading(doc, "九、后续行动项", 1)
actions = [
    "将6.6版总控计划作为本轮客户基准，不再要求客户先补齐全线逐构件明细。",
    "与总工确认关键工点清单、控制节点及其最迟完成时间，重点核对架梁和合龙约束。",
    "选择永宁河特大桥主桥或两河口大桥作为控制性工程样例，整理最小精排数据集。",
    "选择一座普通桥梁形成聚合计划，记录与旧版逐墩方式在对象数量和维护步骤上的差异。",
    "以9月1日开工条件完成一次局部重算和节点影响分析。",
    "组织总工、工程部长复核，收集对节点可信度、维护成本和采用意愿的明确反馈。",
]
for item in actions:
    add_list_item(doc, item, decimal_id)

add_heading(doc, "十、风险与边界", 1)
for item in [
    "本轮已验证客户采用分层总控方式管理节点，但尚未验证当前产品可以无损承接和持续维护该计划。",
    "项目尚未正式开工，月度进度反馈、滚动重排和风险闭环仍缺少真实执行数据。",
    "新版中仍有多项计划从7月1日开始，而现场最新口径为9月1日开工；需由客户确认哪些属于筹备工作、哪些需要整体顺延。",
    "新版未检出公式缓存错误，只能说明保存结果未见明显错误，不能替代客户对计算逻辑和节点日期的专业复核。",
    "控制性工程清单、关键节点及预警余量仍需由总工确认，不能仅由系统自行判断。",
    "当前结论不等于放弃逐构件数据；明细数据仍可用于计算、追溯和关键工程分析，但不应全部转化为用户必须维护的计划任务。",
]:
    add_list_item(doc, item, bullet_id)

add_heading(doc, "附：资料口径", 1)
add_body(doc, "源文件包括《泸古高速TJ-1标桥梁进度统计表4.27(1).xlsx》和《泸古高速项目TJ-1标总体进度计划-1标6.6.xlsx》，两份文件均未做修改。工作簿规模由程序读取工作表、单元格、公式和字段统计得到。")
add_body(doc, "客户已明确按6.6版总控计划管控节点；该事实用于确认计划管理粒度。关于系统导入、滚动调整、风险预警和持续使用价值的判断，仍属于下一阶段待验证产品假设。", italic=True)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
