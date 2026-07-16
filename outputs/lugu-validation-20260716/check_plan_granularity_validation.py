from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
PATH = ROOT / "docs" / "泸古项目计划粒度验证记录_20260716_v2.docx"
OLD_SOURCE = Path(
    r"D:\00-1-03-生产产线-24年\02-产品需求\15-2026斑马基建版\03 项目验证\泸古1标\泸古高速TJ-1标桥梁进度统计表4.27(1).xlsx"
)
NEW_SOURCE = Path(r"C:\Users\wangl-bj\Desktop\泸古高速项目TJ-1标总体进度计划-1标6.6.xlsx")

with ZipFile(PATH) as archive:
    assert archive.testzip() is None

doc = Document(PATH)
assert len(doc.sections) == 1
section = doc.sections[0]
assert round(section.page_width.inches, 2) == 8.50
assert round(section.page_height.inches, 2) == 11.00
assert all(
    round(value.inches, 2) == 1.00
    for value in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin)
)

paragraphs = [p for p in doc.paragraphs if p.text.strip()]
headings = [p for p in paragraphs if p.style.name.startswith("Heading")]
numbered = [p for p in paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
all_text = "\n".join(
    [p.text for p in paragraphs]
    + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
)

assert paragraphs[0].text == "泸古项目计划粒度验证记录"
assert len(headings) >= 15
assert len(numbered) >= 30
assert len(doc.tables) == 5
assert "客户已明确将6.6版作为节点管控计划" in all_text
assert "减少 73.3%" in all_text
assert "项目总控＋控制性工程精排＋普通工程形象进度" in all_text
assert "PLACEHOLDER" not in all_text
assert "说话人1" not in all_text
assert OLD_SOURCE.exists()
assert NEW_SOURCE.exists()

for table in doc.tables:
    tbl_pr = table._tbl.tblPr
    widths = [int(x.get(qn("w:w"))) for x in table._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert sum(widths) == 9360, widths
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    layout = tbl_pr.find(qn("w:tblLayout"))
    assert tbl_w is not None and tbl_w.get(qn("w:w")) == "9360"
    assert tbl_ind is not None and tbl_ind.get(qn("w:w")) == "120"
    assert layout is not None and layout.get(qn("w:type")) == "fixed"
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            assert tc_w is not None and int(tc_w.get(qn("w:w"))) == widths[idx]
            assert cell.text.strip()

print(
    {
        "path": str(PATH),
        "bytes": PATH.stat().st_size,
        "paragraphs": len(paragraphs),
        "headings": len(headings),
        "real_lists": len(numbered),
        "tables": len(doc.tables),
        "zip_ok": True,
        "geometry_ok": True,
        "old_source_unchanged_bytes": OLD_SOURCE.stat().st_size,
        "new_source_unchanged_bytes": NEW_SOURCE.stat().st_size,
    }
)
