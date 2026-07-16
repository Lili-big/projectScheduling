from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


INPUT = Path(
    r"D:\00-1-03-生产产线-24年\02-产品需求\15-2026斑马基建版\03 项目验证\泸古1标\泸古高速TJ-1标桥梁进度统计表4.27(1).xlsx"
)
OUTPUT = Path(__file__).with_name("bridge_schedule_structure.json")


def non_empty_count(ws):
    return sum(1 for row in ws.iter_rows() for cell in row if cell.value is not None)


def sheet_profile(ws, values_ws):
    formulas = []
    formula_columns = set()
    date_cells = 0
    errors = []
    pier_labels = set()
    pile_ids = set()
    resource_cells = 0
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if cell.data_type == "f":
                formulas.append((cell.coordinate, str(value)))
                formula_columns.add(cell.column_letter)
            if hasattr(value, "year") and hasattr(value, "month"):
                date_cells += 1
            if isinstance(value, str):
                if re.search(r"\d+[a-zA-Z]?\s*#?\s*(墩|台)", value):
                    pier_labels.add(value.strip())
                if re.fullmatch(r"\d+[a-zA-Z]?-\d+", value.strip()):
                    pile_ids.add(value.strip())
                if any(term in value for term in ("旋挖钻", "冲击钻", "模板", "塔吊", "电梯", "挂篮")):
                    resource_cells += 1
    for row in values_ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("#"):
                errors.append((cell.coordinate, cell.value))
    sample_formulas = [
        {"cell": coordinate, "formula": formula}
        for coordinate, formula in formulas[:8]
    ]
    return {
        "name": ws.title,
        "state": ws.sheet_state,
        "rows": ws.max_row,
        "columns": ws.max_column,
        "non_empty_cells": non_empty_count(ws),
        "merged_ranges": None,
        "formula_count": len(formulas),
        "formula_columns": sorted(formula_columns),
        "date_cells": date_cells,
        "cached_errors": errors[:20],
        "cached_error_count": len(errors),
        "pier_label_count": len(pier_labels),
        "pile_id_count": len(pile_ids),
        "resource_text_cells": resource_cells,
        "sample_formulas": sample_formulas,
    }


wb = load_workbook(INPUT, data_only=False, read_only=True)
values_wb = load_workbook(INPUT, data_only=True, read_only=True)
profiles = [sheet_profile(ws, values_wb[ws.title]) for ws in wb.worksheets]
detail = [item for item in profiles if item["name"] != "汇总表"]
summary = {
    "file": str(INPUT),
    "sheet_count": len(profiles),
    "visible_sheet_count": sum(item["state"] == "visible" for item in profiles),
    "sheet_names": [item["name"] for item in profiles],
    "detail_sheet_count": len(detail),
    "total_rows_across_sheets": sum(item["rows"] for item in profiles),
    "total_non_empty_cells": sum(item["non_empty_cells"] for item in profiles),
    "total_formulas": sum(item["formula_count"] for item in profiles),
    "total_date_cells": sum(item["date_cells"] for item in profiles),
    "total_pile_ids": sum(item["pile_id_count"] for item in detail),
    "total_pier_labels": sum(item["pier_label_count"] for item in detail),
    "total_resource_text_cells": sum(item["resource_text_cells"] for item in detail),
    "cached_formula_error_count": sum(item["cached_error_count"] for item in profiles),
    "detail_rows_min": min(item["rows"] for item in detail),
    "detail_rows_max": max(item["rows"] for item in detail),
    "detail_cols_min": min(item["columns"] for item in detail),
    "detail_cols_max": max(item["columns"] for item in detail),
}
OUTPUT.write_text(json.dumps({"summary": summary, "sheets": profiles}, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
print("DETAIL_SHEETS")
for item in detail:
    print(
        item["name"],
        f"{item['rows']}x{item['columns']}",
        f"formulas={item['formula_count']}",
        f"piles={item['pile_id_count']}",
        f"piers={item['pier_label_count']}",
        f"errors={item['cached_error_count']}",
    )
print("SAMPLE_FORMULAS")
for item in detail[:2]:
    print(item["name"], item["sample_formulas"])
