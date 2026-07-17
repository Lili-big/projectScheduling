from __future__ import annotations

import importlib.util
from pathlib import Path

from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
BASE_BUILDER = Path(__file__).with_name("build_plan_granularity_validation.py")
OUTPUT = ROOT / "docs" / "泸古项目计划管理方式变化验证记录_20260716_v3.docx"


spec = importlib.util.spec_from_file_location("lugu_validation_base", BASE_BUILDER)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)

doc = base.doc
body = doc._element.body
for child in list(body):
    if not child.tag.endswith("sectPr"):
        body.remove(child)


doc.add_paragraph("泸古项目计划管理方式变化验证记录", style="Title")
doc.add_paragraph("从逐桥精细排程到总控节点管理的业务验证｜V3", style="Subtitle")

metadata = [
    ("项目", "泸古高速 TJ-1 标"),
    ("客户角色", "工程部长"),
    ("验证材料", "前期桥梁明细计划、6.6 版总控计划及客户访谈反馈"),
    ("客户当前口径", "以 6.6 版总控计划管控项目节点"),
    ("项目阶段", "开工前临建阶段，当前计划 9 月 1 日开工"),
    ("验证状态", "计划管理方向已确认，产品价值闭环待验证"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    base.set_run_font(p.add_run(f"{label}："), size=10.5, bold=True, color="000000")
    base.set_run_font(p.add_run(value), size=10.5, color="000000")

rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(6)
rule.paragraph_format.space_after = Pt(10)
base.add_paragraph_border(rule, bottom={"color": base.COLORS["blue"], "size": 14, "space": 2})

base.add_callout(
    doc,
    "核心结论",
    "客户不是不需要工期计算，而是不再接受把所有桥梁、墩台和工序都变成长期维护的计划任务。当前做法是先用总控计划管住专业、工点和关键节点，只对会影响架梁、合龙和总体完工的控制性工程继续深入。产品应承接这一管理方式，而不是要求客户重新回到全量精细排程。",
    tone="green",
)


base.add_heading(doc, "一、总体判断", 1)
for item in [
    "用户的问题真实存在。前期逐桥、逐墩编排虽然看起来完整，但编制、维护和调整成本过高，无法成为日常计划管理方式。",
    "客户已经形成新的管理选择：以总控计划管节点，普通工程保持必要的管理粒度，控制性工程再按关键墩位和工序展开。",
    "这次变化不是简单删减计划内容，而是管理目标从“把所有工作排清楚”转向“保证关键节点不失控”。",
    "产品机会已经从“替客户重新生成一套详细计划”转向“承接现有总控计划、识别关键影响、支持月度滚动和节点预警”。",
    "当前只验证了方向，尚未验证系统能否比 Excel 更省维护、更快调整，并形成客户愿意持续使用的闭环。",
]:
    base.add_list_item(doc, item, base.bullet_id)


base.add_heading(doc, "二、用户之前怎么做", 1)
base.add_heading(doc, "2.1 原来的计划思路", 2)
base.add_body(doc, "客户前期希望把每座桥的施工过程完整算清楚。做法是先按桥梁工点逐个拆分，再把桩基、承台、系梁、墩柱、盖梁等施工对象继续往下细化，结合工程量、工效、设备、模板和施工顺序计算持续时间，最终形成每项工作的开始和结束安排。")
base.add_body(doc, "这套思路的出发点是合理的：在项目策划阶段尽可能把工期、资源和施工组织想清楚，提前判断每座桥什么时候能完成、需要配置多少资源、是否会影响架梁和总体工期。")

base.add_heading(doc, "2.2 实际使用中出现的问题", 2)
for item in [
    "前期编制量过大。为了形成完整计划，需要持续补充大量工程量、工效、资源和施工关系。",
    "现场一旦发生开工条件、资源或施工顺序变化，后续大量任务都要跟着调整，计划很快失去可维护性。",
    "普通桥梁现场不会严格按照每个墩台的预定日期组织施工，班组往往根据工作面和现场条件灵活安排。",
    "控制性工程和普通工程被用同一种粒度管理，投入了大量维护成本，却没有产生同等的管理收益。",
    "计划逐渐变成一套复杂的计算底表，而不是工程部长能够持续用来下达任务、检查偏差和判断节点的管理工具。",
]:
    base.add_list_item(doc, item, base.bullet_id)

base.add_callout(
    doc,
    "原方式的核心问题",
    "计划的详细程度超过了现场实际控制程度。问题不在于算得不够准，而在于要求用户维护了太多并不需要逐项管理的对象。",
    tone="yellow",
)


base.add_heading(doc, "三、用户现在怎么做", 1)
base.add_body(doc, "客户现在以 6.6 版总控计划作为节点管控依据。计划首先从全标段视角安排路基、桥梁、隧道、互通、梁场和架梁等专业及工点，再区分普通工程与控制性工程，采用不同的管理深度。")

for item in [
    "第一层：管全标段总控。关注专业、主要工点、梁场、架梁路线和总体完工等关键节点。",
    "第二层：管控制性工程。对永宁河特大桥主桥、两河口大桥等关键工程，继续细化到主墩、挂篮、悬浇和合龙等决定工期的过程。",
    "第三层：管普通工程。普通桥梁主要按下部结构、左右幅架梁和桥面系等阶段管理，后续执行更适合转成月度形象进度量。",
    "计算依据继续保留。工程量、工效、资源配置和施工顺序仍用于判断工期，但不再要求全部变成用户每天维护的明细任务。",
]:
    base.add_list_item(doc, item, base.decimal_id)

base.add_heading(doc, "3.1 当前管理逻辑", 2)
base.add_body(doc, "客户当前的管理链条可以概括为：先确定全标段总控节点，再识别真正影响节点的控制性工程；关键工程深入管过程，普通工程主要管阶段完成和形象进度；发生变化时，先判断是否侵蚀架梁、合龙或完工节点，再决定是否调整后续计划。")


base.add_heading(doc, "四、为什么改成现在的方式", 1)
base.add_heading(doc, "4.1 管理目标变了", 2)
base.add_body(doc, "工程部长真正需要的不是一张最细的计划表，而是一套能够回答管理问题的计划：哪些节点必须守住、当前工作是否会影响这些节点、需要提前安排什么资源、发生偏差后应该调整哪里。")

base.add_heading(doc, "4.2 计划必须跟现场的控制方式一致", 2)
base.add_body(doc, "控制性工程决定架梁、合龙和总体工期，值得投入精力细管；普通工程的具体施工顺序受工作面、班组、设备、天气和临时条件影响较大，更适合控制阶段目标和完成量。不同工程使用不同粒度，反而更接近真实施工管理。")

base.add_heading(doc, "4.3 计划必须长期维护得动", 2)
base.add_body(doc, "项目当前还在临建阶段，开工时间和施工条件仍可能变化。此时把所有工作排得过细，计划会在真正开工前反复失效。先建立稳定的总控框架，再随着条件明确逐步展开，维护成本更可控。")

base.add_heading(doc, "4.4 调整应围绕影响，而不是全盘重做", 2)
base.add_body(doc, "客户更关心一项变化会不会影响后续架梁、合龙或完工节点。只有影响真正传导到控制节点时，才有必要进一步调资源、改顺序或压缩工期，而不是每次变化都重排所有任务。")


base.add_heading(doc, "五、前后管理方式的业务差异", 1)
base.add_table(
    doc,
    ["业务维度", "之前的做法", "现在的做法", "变化的意义"],
    [
        ["管理目标", "把所有工作排细、算清", "守住全标段关键节点", "从计划完整转向结果可控"],
        ["控制对象", "桥梁、墩台和具体工序", "专业、工点、关键节点", "减少无效维护对象"],
        ["计划粒度", "全项目统一精细化", "控制工程细、普通工程粗", "计划粒度匹配管理粒度"],
        ["普通工程", "逐墩逐项给日期", "按桥、幅、阶段和完成量", "给现场保留执行弹性"],
        ["计划调整", "变化后大量任务联动修改", "先看节点影响，再局部调整", "调整围绕决策而不是表格"],
        ["计划价值", "形成完整计算结果", "支持下达、检查、预警和决策", "从计算工具转为管理工具"],
    ],
    [1500, 2500, 2500, 2860],
    compact=True,
)


base.add_heading(doc, "六、当前产品与客户方式的差距", 1)
base.add_heading(doc, "6.1 进入方式不一致", 2)
base.add_body(doc, "当前 Demo 更强调从项目结构和参数出发生成计划，而客户已经完成总控计划。客户下一步更需要系统先读懂并承接现有计划，再在此基础上增强，而不是重新要求他从头编制。")

base.add_heading(doc, "6.2 计划粒度不一致", 2)
base.add_body(doc, "当前能力更容易形成统一的明细任务网络，客户则需要同一项目中同时存在总控、专业、普通桥梁和控制性工程等不同粒度，并能够按角色展开或聚合。")

base.add_heading(doc, "6.3 动态调整闭环尚未形成", 2)
base.add_body(doc, "客户关心的是月度计划下达、实际完成反馈、未完成原因、剩余工期判断和后续计划调整。当前 Demo 主要验证了计划生成和方案计算，还没有用真实进度跑通一次滚动管理。")

base.add_heading(doc, "6.4 预警口径还需要收敛", 2)
base.add_body(doc, "客户不需要所有任务一有日期偏差就报警，而是希望系统判断偏差是否正在侵蚀架梁、合龙和总体完工等关键节点余量。预警必须围绕管理后果，而不是简单比较计划日期。")

base.add_heading(doc, "6.5 产品价值尚未被实际使用验证", 2)
base.add_body(doc, "项目尚未正式开工，客户还没有在系统中完成真实数据更新、计划调整和结果复核。因此，方向匹配不等于价值闭环已经成立。")


base.add_heading(doc, "七、对产品验证的启示", 1)
base.add_heading(doc, "7.1 已经可以确认的方向", 2)
for item in [
    "产品应以客户现有总控计划为基准，而不是以重新生成全量计划为前提。",
    "计划必须支持分层管理：全标段总控、控制性工程精排、普通工程阶段和形象进度管理。",
    "精细计算仍有价值，但计算粒度与用户维护粒度必须分开；系统可以算细，用户不必管细。",
    "动态调整和风险预警必须围绕关键节点影响，避免把普通任务偏差放大成管理噪声。",
]:
    base.add_list_item(doc, item, base.bullet_id)

base.add_heading(doc, "7.2 不需要继续大范围验证的方向", 2)
base.add_body(doc, "不需要再投入大量时间证明“系统能否把所有桥梁都拆细并自动排出一套计划”。客户已经用实际选择说明，这不是他愿意长期维护的方式。前期策划计算仍需验证，但应集中在控制性工程和真实决策场景上。")

base.add_heading(doc, "7.3 接下来真正需要验证的价值", 2)
for item in [
    "能否快速承接并准确还原客户当前总控计划。",
    "能否在不增加全量明细维护的情况下，对控制性工程给出可信工期判断。",
    "条件变化后，能否快速指出受影响的工点、架梁路径和关键节点。",
    "能否从总控计划形成可执行的月度计划，并根据实际完成情况滚动更新。",
    "客户是否认为系统比现有 Excel 更容易维护，并愿意持续更新和复核。",
]:
    base.add_list_item(doc, item, base.bullet_id)

base.add_callout(
    doc,
    "产品验证焦点",
    "从“证明系统会排计划”转向“证明系统能把客户现有总控计划管起来”：接得住、看得懂、调得动、能预警，并且客户维护得动。",
    tone="blue",
)


base.add_heading(doc, "八、下一步验证安排", 1)
base.add_heading(doc, "8.1 开工前：验证计划承接与一次调整", 2)
for item in [
    "以 6.6 版总控计划作为基准，在系统中还原专业、工点、关键节点、梁场和架梁顺序。",
    "选择一座控制性桥梁，复现其关键墩位和工序逻辑，由总工或工程部长复核计算差异。",
    "选择一座普通桥梁，按桥、左右幅和施工阶段形成聚合计划，确认这种粒度是否便于管理。",
    "以 9 月 1 日开工为变化条件，完成一次局部调整，说明哪些节点受影响、为什么、建议如何处理。",
    "让客户评价三件事：结果是否可信、调整是否省事、是否愿意继续用真实数据验证。",
]:
    base.add_list_item(doc, item, base.decimal_id)

base.add_heading(doc, "8.2 开工后：验证一个月度滚动闭环", 2)
for item in [
    "从总控计划生成一个真实月度计划。",
    "录入一次实际完成量、未完成任务和原因。",
    "判断偏差是否影响控制节点，并更新剩余计划。",
    "对比不调整与采取措施后的节点结果。",
    "由客户确认调整方案，并评价是否愿意按月持续使用。",
]:
    base.add_list_item(doc, item, base.decimal_id)

base.add_heading(doc, "8.3 本轮验证成功的判断标准", 2)
for item in [
    "客户现有总控计划能够被系统准确承接，不要求重新整理一套全量明细。",
    "控制性工程的关键节点差异能够解释清楚，并获得客户专业复核。",
    "普通工程的计划粒度符合现场管理习惯，维护工作量明显可接受。",
    "开工变化后，系统能够指出真正受影响的节点，而不是全盘重排和大量误报。",
    "客户愿意提供下一轮真实数据，并明确复核人和复核时间。",
]:
    base.add_list_item(doc, item, base.bullet_id)


base.add_heading(doc, "九、验证边界", 1)
base.add_body(doc, "本轮已经确认的是客户的计划管理方式：以总控计划管节点，对控制性工程深入、普通工程聚合。尚未确认的是当前产品能否承接该方式并产生持续价值。")
base.add_body(doc, "项目尚未正式开工，月度进度反馈、滚动调整、节点预警和持续使用意愿仍需等待真实施工数据验证。客户愿意沟通和提供资料，不等同于已经承诺正式试点或生产使用。")
base.add_body(doc, "本报告用于内部产品验证判断。客户业务事实、产品推断和待验证假设已分别表达，不将产品设想写成客户已经认可的能力。", italic=True)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
